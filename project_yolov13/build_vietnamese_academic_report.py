from __future__ import annotations

import csv
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "results"
BENCH = list(csv.DictReader((OUT / "adv_tshirt_patch_benchmark.csv").open(encoding="utf-8-sig", newline="")))
ABL = list(csv.DictReader((OUT / "adv_tshirt_patch_ablation.csv").open(encoding="utf-8-sig", newline="")))
abmap = {x["variant"]: x for x in ABL}
MD = OUT / "BaoCaoKhoaHoc_YOLOv13_AdvPatch.md"
DOCX = OUT / "BaoCaoKhoaHoc_YOLOv13_AdvPatch.docx"
families = {"YOLOv8": 0, "YOLOv10": 1, "YOLO11": 2, "YOLOv12": 3, "YOLOv13": 4}
scales = {"n": 0, "s": 1, "l": 2}

def row(model, split, device):
    return next(x for x in BENCH if x["model_name"] == model and x["split"] == split and x["device"] == device)

def num(x, digits=2):
    return f"{float(x):.{digits}f}"

def signed(x, digits=2):
    value = float(x)
    if abs(value) < 0.5 * 10 ** (-digits):
        value = 0.0
    return f"{value:+.{digits}f}"

def p(row, key):
    return float(row[key]) * 100

test = [x for x in BENCH if x["split"] == "test" and x["device"] == "GPU"]
test.sort(key=lambda x: (scales[x["scale"]], families[x["family"]]))
head = ["Mô hình","P (%)","R (%)","mAP@0.5 (%)","mAP@0.5:0.95 (%)","Params (M)","FLOPs (G)","CPU (ms)","CPU FPS*","GPU (ms)","GPU FPS*"]
bm = ["| " + " | ".join(head) + " |", "|" + "|".join(["---"]*len(head)) + "|"]
latex = []
for x in test:
    c = row(x["model_name"], "test", "CPU")
    cells = [x["model_name"],num(p(x,"precision")),num(p(x,"recall")),num(p(x,"map50")),num(p(x,"map50_95")),num(x["params_m"]),num(x["flops_g"]),num(c["latency_ms"]),num(c["fps"]),num(x["latency_ms"]),num(x["fps"])]
    bm.append("| " + " | ".join(cells) + " |")
    latex.append(" & ".join([cells[0],*cells[1:]]) + r" \\")
val = [x for x in BENCH if x["split"] == "val" and x["device"] == "GPU"]
th = ["| Quy mô | Baseline tốt nhất theo validation | Val mAP95 (%) | Baseline test mAP95 (%) | YOLOv13 | YOLOv13 val mAP95 (%) | YOLOv13 test mAP95 (%) | Δ test (điểm %) | Baseline GPU (ms) | YOLOv13 GPU (ms) | Δ latency (ms) |","|"+"|".join(["---"]*11)+"|"]
trade_points = []
for scale in ("n","s","l"):
    cand = [x for x in val if x["scale"] == scale]
    base = max(cand, key=lambda x: float(x["map50_95"]))
    y13 = next(x for x in cand if x["family"] == "YOLOv13")
    bt = row(base["model_name"],"test","GPU")
    yt = row(y13["model_name"],"test","GPU")
    label = {"n":"Nano","s":"Small","l":"Large"}[scale]
    vals = [label,base["model_name"],num(p(base,"map50_95")),num(p(bt,"map50_95")),y13["model_name"],num(p(y13,"map50_95")),num(p(yt,"map50_95")),signed(p(yt,"map50_95")-p(bt,"map50_95")),num(bt["latency_ms"]),num(yt["latency_ms"]),signed(float(yt["latency_ms"])-float(bt["latency_ms"]),2)]
    th.append("| "+" | ".join(vals)+" |")
    trade_points.extend([(scale, "baseline", base, bt), (scale, "YOLOv13", y13, yt)])

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.lines import Line2D

PLOTS = OUT / "adv_tshirt_patch_results"
PLOTS.mkdir(parents=True, exist_ok=True)
markers = {"n": "o", "s": "s", "l": "^"}
colors = {"baseline": "#3465A4", "YOLOv13": "#C43C39"}
offsets = {
    ("n", "baseline"): (-12, 12), ("n", "YOLOv13"): (10, -10),
    ("s", "baseline"): (-48, 16), ("s", "YOLOv13"): (10, 16),
    ("l", "baseline"): (-48, -18), ("l", "YOLOv13"): (10, 10),
}

def make_tradeoff(xkey, xlabel, title, filename):
    fig, ax = plt.subplots(figsize=(7.6, 4.25), constrained_layout=True)
    for scale, kind, valrow, testrow in trade_points:
        x = float(testrow[xkey])
        y = p(testrow, "map50_95")
        ax.scatter(x, y, s=58, marker=markers[scale], color=colors[kind], zorder=3)
        ax.annotate(testrow["model_name"], (x, y), xytext=offsets[(scale, kind)],
                    textcoords="offset points", fontsize=8,
                    arrowprops={"arrowstyle": "-", "color": "#777777", "lw": 0.55})
    ax.set_xlabel(xlabel, fontsize=9)
    ax.set_ylabel("mAP@0.5:0.95 test (%)", fontsize=9)
    ax.set_title(title, fontsize=11)
    ax.grid(True, alpha=0.25)
    ax.set_ylim(78.0, 85.2)
    legend = [Line2D([0],[0], marker="o", color="w", markerfacecolor=colors[k], label=k, markersize=7)
              for k in ("baseline", "YOLOv13")]
    ax.legend(handles=legend, frameon=False, fontsize=8, loc="lower right")
    fig.savefig(PLOTS / filename, dpi=200, facecolor="white")
    plt.close(fig)

make_tradeoff("latency_ms", "Độ trễ GPU (ms/ảnh, batch 4)",
              "Đánh đổi độ chính xác và độ trễ trên ba quy mô",
              "advpatch_tradeoff_selected_latency.png")
make_tradeoff("flops_g", "FLOPs tại 416×416 (G)",
              "Đánh đổi độ chính xác và chi phí tính toán",
              "advpatch_tradeoff_selected_flops.png")

fig, axes = plt.subplots(1, 2, figsize=(9.4, 3.35), constrained_layout=True)
ablation_labels = ["Đầy đủ", "Bỏ HyperACE", "Bỏ FullPAD", "Thay DS-C3k2"]
ablation_map = [p(abmap[k], "map50_95") for k in ("YOLOv13-n-full", "YOLOv13-n-no-hyperace", "YOLOv13-n-no-fullpad", "YOLOv13-n-no-dsc3k2")]
ablation_latency = [float(abmap[k]["latency_ms"]) for k in ("YOLOv13-n-full", "YOLOv13-n-no-hyperace", "YOLOv13-n-no-fullpad", "YOLOv13-n-no-dsc3k2")]
bar_colors = ["#3465A4", "#888A85", "#888A85", "#888A85"]
for ax, values, title, ylabel, ylim in (
    (axes[0], ablation_map, "Độ chính xác", "mAP@0.5:0.95 (%)", (78, 85)),
    (axes[1], ablation_latency, "Độ trễ", "GPU latency (ms)", (4.5, 7.0)),
):
    bars = ax.bar(ablation_labels, values, color=bar_colors)
    ax.set_title(title, fontsize=10)
    ax.set_ylabel(ylabel, fontsize=8)
    ax.set_ylim(*ylim)
    ax.tick_params(axis="x", labelrotation=18, labelsize=7)
    ax.tick_params(axis="y", labelsize=8)
    ax.grid(axis="y", alpha=0.22)
    ax.set_axisbelow(True)
    for bar, value in zip(bars, values):
        ax.annotate(f"{value:.2f}", (bar.get_x()+bar.get_width()/2, value),
                    xytext=(0, 3), textcoords="offset points", ha="center", fontsize=7)
fig.savefig(PLOTS / "advpatch_ablation_map_latency.png", dpi=200, facecolor="white")
plt.close(fig)

abmap = {x["variant"]:x for x in ABL}
full = abmap["YOLOv13-n-full"]
ah = ["| Biến thể | P (%) | R (%) | mAP@0.5 (%) | mAP@0.5:0.95 (%) | Params (M) | FLOPs (G) | GPU (ms) | Δ mAP95 (điểm %) | Δ Params (M) | Δ FLOPs (G) | Δ latency (ms) |","|"+"|".join(["---"]*12)+"|"]
for name in ("YOLOv13-n-full","YOLOv13-n-no-hyperace","YOLOv13-n-no-fullpad","YOLOv13-n-no-dsc3k2"):
    x=abmap[name]
    vals=[name,num(float(x["precision"])*100),num(float(x["recall"])*100),num(float(x["map50"])*100),num(float(x["map50_95"])*100),num(x["params_m"]),num(x["flops_g"]),num(x["latency_ms"]),signed((float(x["map50_95"])-float(full["map50_95"]))*100),signed(float(x["params_m"])-float(full["params_m"])),signed(float(x["flops_g"])-float(full["flops_g"])),signed(float(x["latency_ms"])-float(full["latency_ms"]))]
    ah.append("| "+" | ".join(vals)+" |")

latex_block = "\n".join([
"    % Requires \\usepackage{booktabs,graphicx}",
"    \\begin{table*}[t]",
"    \\centering",
"    \\scriptsize",
"    \\caption{Kết quả fine-tune trên test split AdvT-shirt-1K dẫn xuất.}",
"    \\label{tab:advtshirt-yolov13}",
"    \\resizebox{\\textwidth}{!}{%",
"    \\begin{tabular}{lrrrrrrrrrr}",
"    \\toprule",
"    Model & P (\\%) & R (\\%) & mAP@0.5 (\\%) & mAP@0.5:0.95 (\\%) & Params (M) & FLOPs (G) & CPU ms & CPU FPS* & GPU ms & GPU FPS* \\\\",
"    \\midrule",
*["    "+x for x in latex],
"    \\bottomrule",
"    \\end{tabular}}",
"    \\end{table*}",
])

front = """# Đánh giá YOLOv13 cho định vị miếng dán đối kháng vật lý trên AdvT-shirt-1K

**Tác giả:** Phạm Tuấn Anh, Mai Quốc Bảo — Trường Đại học Công nghệ TP.HCM (HUTECH)  
**Giảng viên hướng dẫn:** Đinh Phương Nam  
**Ngày báo cáo:** 05 tháng 10 năm 2026

## Tóm tắt

Miếng dán đối kháng in có thể làm suy giảm tính toàn vẹn của bộ phát hiện vật thể khi đưa mô hình từ môi trường số ra thế giới thực. Nghiên cứu này đánh giá YOLOv13 trên tác vụ định vị hộp miếng dán, đồng thời đối sánh có kiểm soát với YOLOv8, YOLOv10, YOLO11 và YOLOv12. Các split có nhãn của bộ dữ liệu NCKH gốc hiện chưa có ảnh và nhãn, vì vậy AdvT-shirt-1K được dùng như tập proxy bên ngoài. Bản mirror có 1.123 ảnh và 1.548 hộp miếng dán; split xác định theo seed 42 gồm 785/169/169 ảnh train/validation/test. Mười lăm mô hình Nano/Small/Large được fine-tune với đầu vào 416×416 và lịch huấn luyện đồng nhất. YOLOv13s đạt 83,96% mAP@0.5:0.95 và 99,45% mAP@0.5 trên test, với 9,55 triệu tham số, 9,15 GFLOPs, độ trễ GPU 3,16 ms/ảnh và CPU 38,95 ms/ảnh. YOLOv12s đạt 83,81% và 99,41%, với độ trễ GPU 2,91 ms/ảnh; đây là cấu hình có mAP@0.5:0.95 validation cao nhất trong 15 cấu hình và do đó được chọn báo cáo chính. YOLOv13s cao hơn YOLOv12s 0,16 điểm phần trăm mAP@0.5:0.95 trên test nhưng tăng 0,25 ms độ trễ GPU. Kết quả chỉ chứng minh khả năng định vị trên proxy, không trực tiếp chứng minh khả năng chống tấn công vật lý, quá tải NMS hay hiệu quả trên dữ liệu đề tài gốc.

**Từ khóa:** phát hiện vật thể; miếng dán đối kháng vật lý; định vị patch; YOLOv13; AdvT-shirt-1K; robustness.

> **Đối chiếu nguồn số liệu:** cặp 83,81% mAP@0.5:0.95 và 99,41% mAP@0.5 thuộc YOLOv12s. YOLOv13s có kết quả tương ứng là 83,96% và 99,45% trong benchmark CSV.

## 1. Giới thiệu và mục tiêu nghiên cứu

Miếng dán đối kháng vật lý được in và đặt trong cảnh thật để gây nhầm lẫn cho bộ phát hiện. Đây là rủi ro đối với tính toàn vẹn của nhận thức thị giác trong phương tiện tự hành và thiết bị biên. Định vị vùng patch có thể hỗ trợ hệ thống nhận diện khu vực nghi vấn và kích hoạt kiểm tra bổ sung. Tuy nhiên, phát hiện được hộp patch không đồng nghĩa mô hình đã chống được tấn công hoặc ngăn được suy giảm dịch vụ. [1], [2]

YOLOv13 đề xuất HyperACE để khai thác tương quan đặc trưng bậc cao giữa nhiều vị trí và tỉ lệ; FullPAD phân phối đặc trưng đã tăng cường xuyên suốt pipeline; DS-C3k2 dùng các khối tích chập tách biệt theo chiều sâu nhằm giảm chi phí tính toán. Đây là giả thuyết kiến trúc từ công trình YOLOv13. Hiệu quả trên AdvT-shirt-1K phải được xác định bằng phép đo tại chỗ, không suy diễn trực tiếp từ kết quả COCO. [3]

Nghiên cứu trả lời ba câu hỏi:

- **RQ1 — Chất lượng định vị:** YOLOv13 định vị hộp miếng dán ra sao trên split AdvT-shirt-1K được chia lại?
- **RQ2 — Chi phí suy luận:** độ chính xác đi kèm độ trễ, FPS, tham số và FLOPs nào so với baseline cùng quy mô?
- **RQ3 — Mức độ phòng thủ:** các phép đo hiện có có chứng minh khả năng chống tấn công vật lý hoặc DoS/NMS-overload trong bài toán NCKH gốc không?

## 2. Thiết lập thực nghiệm

### 2.1. Dữ liệu và nhãn

AdvT-shirt-1K gồm ảnh áo có miếng dán trong bối cảnh vật lý, với nhãn hộp cho người và patch. Nghiên cứu này chỉ dùng nhãn patch và một lớp duy nhất. Bản mirror hiện có 1.123 ảnh, thấp hơn con số 1.131 ảnh mà nguồn tác giả nêu; chênh lệch tám ảnh chưa được giải quyết và được xem là hạn chế về tính toàn vẹn nguồn dữ liệu. [1], [2]

Nguồn mirror không cung cấp split phù hợp cho quy trình này. Nhóm tạo split xác định 70/15/15 theo seed 42, phân tầng trong năm thư mục nguồn: train 785 ảnh/1.106 hộp; validation 169 ảnh/220 hộp; test 169 ảnh/222 hộp. Tổng cộng 1.548 hộp. Split được tạo ở mức ảnh, không theo người mặc, thiết kế patch hoặc phiên chụp; vì vậy ảnh tương quan có thể xuất hiện ở nhiều split. Hai nhãn lệch biên 5×10⁻⁷ do làm tròn đã được chặn về [0,1] trong bản split dẫn xuất.

Ảnh gốc rộng 856–2056 pixel và cao 1044–1598 pixel; mọi mô hình nhận đầu vào 416×416. Dataset YAML proxy là dataset/adv_tshirt_patch_detection.yaml. Cấu hình chuẩn dataset/data.yaml dùng path: ./dataset; các split train/validation/test của đề tài hiện có 0 ảnh và 0 nhãn, tên lớp object vẫn là placeholder và không được dùng để huấn luyện.

### 2.2. Phần cứng và phần mềm

| Thành phần | Cấu hình |
|---|---|
| CPU | Intel Core i5-14600KF |
| GPU | NVIDIA GeForce RTX 5060, 7,96 GiB VRAM |
| RAM | 31,79 GiB |
| CUDA | Runtime 12.8 |
| Python | 3.14.7 |
| PyTorch | 2.11.0+cu128 |
| Ultralytics | 8.3.63 |
| Môi trường | project_yolov13/.venv |
| Cấu hình YOLO | project_yolov13/.config |
| Dữ liệu, runs, weights | project_yolov13/dataset, runs, weights |

Các đường dẫn dữ liệu, cache tải dữ liệu, cấu hình YOLO, checkpoint, log và kết quả nằm dưới thư mục dự án. Cây dataset được kiểm tra không có symlink/junction. Không có phép đo trên Jetson hay thiết bị edge chuyên dụng.

### 2.3. Cấu hình huấn luyện

| Siêu tham số | Giá trị |
|---|---|
| Kích thước đầu vào | 416×416 |
| Batch size | 4 |
| Epoch tối đa / patience | 50 / 10 |
| Optimizer | SGD |
| Learning rate ban đầu | 0,01 |
| Scheduler | tuyến tính, lr cuối bằng 0,01 lần lr đầu |
| Warm-up | 2 epoch |
| Weight decay | 0,0005 |
| Seed / deterministic | 42 / bật |
| HSV | h=0,015; s=0,7; v=0,4 |
| Hình học | translate=0,1; scale=0,5; horizontal flip=0,5 |
| Mosaic | 1,0 |
| MixUp / CopyPaste | 0 / 0 |

Mười lăm cấu hình dùng cùng split, seed, pretrained initialization tương ứng, augmentation và lịch huấn luyện. Mỗi cấu hình chỉ chạy một seed nên chưa có ước lượng phương sai hoặc kiểm định ý nghĩa thống kê.

### 2.4. Quy trình chọn mô hình và đo lường

Validation được dùng chọn mô hình/checkpoint; test chỉ dùng báo cáo cuối. YOLOv12s đạt validation mAP@0.5:0.95 cao nhất toàn bộ 15 cấu hình (84,78%), và đạt 83,81% mAP@0.5:0.95 cùng 99,41% mAP@0.5 trên test. Trong so sánh cùng scale, baseline được chọn là model có validation mAP@0.5:0.95 cao nhất ở scale đó.

Latency là thời gian suy luận trung bình mỗi ảnh theo Ultralytics validator tại batch 4, trên CPU/GPU; không gồm thời gian nạp mô hình. FPS* được tính bằng 1000/latency_ms, là nghịch đảo latency bình quân trên ảnh, không phải phép đo throughput video độc lập. Params và FLOPs do profiler Ultralytics đo ở 416 pixel. Các phép đo phụ thuộc thiết bị và runtime, không so trực tiếp với TensorRT FP16 trên Tesla T4.

<!-- LANDSCAPE -->

## 3. Kết quả và benchmark

### 3.1. So sánh trên test split

**Bảng 1.** Mười lăm mô hình fine-tune trên test split AdvT-shirt-1K dẫn xuất. P/R/mAP tính theo phần trăm; latency là ms/ảnh ở batch 4; FPS* suy ra từ latency.
"""
mid = """
### 3.2. Đánh đổi độ chính xác và tốc độ

**Bảng 2.** YOLOv13 so với baseline cùng scale có validation mAP@0.5:0.95 cao nhất. Test chỉ được dùng sau khi chọn baseline.

"""
trade_text = """
Ở Nano, YOLOv8n được chọn theo validation; YOLOv13n thấp hơn 4,21 điểm phần trăm mAP@0.5:0.95 trên test và có latency GPU cao hơn. Ở Small, YOLOv13s cao hơn YOLOv12s 0,16 điểm phần trăm nhưng chậm hơn 0,25 ms/ảnh (3,16 so với 2,91 ms). Ở Large, YOLOv13l thấp hơn YOLOv12l 0,41 điểm phần trăm và chậm hơn 1,79 ms/ảnh. Benchmark này không cho thấy YOLOv13 vượt trội đồng đều trên mọi scale.

![Hình 1. mAP@0.5:0.95 và độ trễ GPU của baseline được chọn cùng YOLOv13](adv_tshirt_patch_results/advpatch_tradeoff_selected_latency.png)

![Hình 2. mAP@0.5:0.95 và FLOPs của baseline được chọn cùng YOLOv13](adv_tshirt_patch_results/advpatch_tradeoff_selected_flops.png)

<!-- LANDSCAPE -->

### 3.3. Ablation các thành phần kiến trúc

Ablation chạy trên YOLOv13-N với cùng pretrained initialization, split, seed, lịch và augmentation. Bảng dùng metric validation; Δ so với model đầy đủ. Biến thể no-DS-C3k2 thay thế 20 convolution mục tiêu.

**Bảng 3.** Ablation YOLOv13-N trên validation; latency đo trên GPU.

"""
ablation_text = """
Bỏ HyperACE làm mAP@0.5:0.95 giảm 0,39 điểm phần trăm; latency đồng thời tăng 0,77 ms. Bỏ FullPAD làm mAP tăng 2,13 điểm nhưng latency tăng 0,93 ms. Thay DS-C3k2 làm mAP tăng 2,50 điểm, song tham số tăng 2,90 triệu, FLOPs tăng 3,03 G và latency tăng 0,53 ms. Vì vậy, ablation không chứng minh cả ba thành phần đều cải thiện đồng thời độ chính xác và tốc độ trên proxy. Một seed chỉ cung cấp bằng chứng mô tả, chưa hỗ trợ kết luận nhân quả phổ quát.

## 4. Thảo luận

### 4.1. Trả lời câu hỏi nghiên cứu

**RQ1 — Chất lượng định vị.** YOLOv13s đạt 83,96% mAP@0.5:0.95 và 99,45% mAP@0.5 trên test. YOLOv12s là cấu hình cao nhất theo validation trên toàn benchmark, đạt 83,81% mAP@0.5:0.95 trên test. Chênh lệch giữa hai mô hình Small là +0,16 điểm phần trăm cho YOLOv13s trên test.

**RQ2 — Chi phí suy luận.** So với YOLOv12s, YOLOv13s có thêm 0,33 triệu tham số, 0,08 GFLOPs, và chậm hơn 0,25 ms/ảnh GPU, đổi lại mAP@0.5:0.95 cao hơn 0,16 điểm trên test. Đây là đánh đổi nhỏ ở scale Small, không phải cải thiện FPS. Kết quả không thay thế benchmark edge; cần đo trên thiết bị đích, với pipeline tiền/hậu xử lý, stream và công suất thực tế.

**RQ3 — Khả năng phòng thủ.** Thí nghiệm chỉ đánh giá định vị hộp patch. Không đo attack success rate, suy giảm recall của đối tượng, số proposal sau NMS, tải NMS hoặc khả năng khái quát hóa sang patch chưa thấy. Do đó, chưa có bằng chứng rằng YOLOv13 chống được tấn công vật lý hoặc ngăn DoS/NMS-overload.

### 4.2. Ý nghĩa các khối YOLOv13

HyperACE hướng đến liên kết nhiều vị trí đặc trưng cùng lúc thay vì chỉ quan hệ từng cặp; FullPAD truyền thông tin tăng cường qua các giai đoạn; DS-C3k2 nhắm giảm phép tính tích chập. Đây là lý do kiến trúc khiến YOLOv13 đáng khảo sát cho patch có tương quan giữa họa tiết, áo và người. Tuy nhiên, bằng chứng địa phương cho thấy bỏ HyperACE làm giảm nhẹ mAP, còn bỏ FullPAD hoặc thay DS-C3k2 lại tăng mAP trong cấu hình và seed này, đồng thời tăng latency hoặc chi phí. Không thể quy chênh lệch giữa các họ mô hình cho riêng một thành phần vì chúng còn khác backbone/head và pretrained weights.

### 4.3. Tính hợp lệ và giới hạn

- **Nội tại:** cùng split, ảnh vào, optimizer, augmentation, seed và validator tăng khả năng đối sánh; khác biệt pretrained weights và early stopping vẫn có thể ảnh hưởng thứ hạng.
- **Split:** chia theo ảnh trong thư mục nguồn, không theo người/patch/session; ảnh tương quan có thể rơi vào nhiều phân vùng.
- **Bên ngoài:** AdvT-shirt-1K là proxy định vị patch trên áo, không thay thế dataset SpongePatch/Physical-Visual DoS của đề tài.
- **Tốc độ:** latency/FPS phụ thuộc GPU, driver, runtime và batch; FPS* là nghịch đảo latency/ảnh, không phải throughput video.
- **Ablation:** một seed, chưa có khoảng tin cậy; cần chạy nhiều seed và test trên thiết kế/người chụp chưa thấy.
- **Triển khai:** chưa đo năng lượng, camera thật, xử lý đầu-cuối hay thiết bị edge chuyên dụng.

Công trình YOLOv13 công bố kết quả COCO để đánh giá kiến trúc. Các số COCO dùng dữ liệu, nhãn, thiết bị và giao thức khác nên không được ghép với kết quả local thành bảng xếp hạng trực tiếp. [3]

## 5. Kết luận và hướng phát triển

Nghiên cứu fine-tune và benchmark 15 cấu hình YOLOv8/10/11/12/13 Nano/Small/Large cho định vị patch trên split dẫn xuất AdvT-shirt-1K. YOLOv13s đạt 83,96% mAP@0.5:0.95 và 99,45% mAP@0.5; so với YOLOv12s, mức tăng là 0,16 điểm phần trăm mAP@0.5:0.95 cùng 0,25 ms latency GPU tăng thêm. Mô hình có validation mAP cao nhất toàn benchmark là YOLOv12s, với test mAP@0.5:0.95 83,81% và mAP@0.5 99,41%. Vì vậy, cặp số 83,81%/99,41% phải được gắn với YOLOv12s.

Kết luận chỉ áp dụng cho định vị patch trên proxy; thí nghiệm chưa xác lập tính kháng đối kháng, hiệu quả ngăn quá tải NMS hoặc lợi thế trên dữ liệu NCKH gốc. Bước tiếp theo là hoàn thiện nhãn đề tài, chia test độc lập theo người/thiết kế patch, lặp nhiều seed, đo attack success và tải NMS trực tiếp, sau đó đánh giá end-to-end trên phần cứng edge.

## Tài liệu tham khảo

[1] AdvT-shirt-1K, dataset card và giấy phép Apache-2.0. [Hugging Face](https://huggingface.co/datasets/xingjunm/Adv-T-shirt).

[2] W. Wang et al., “BadPatch: Diffusion-Based Generation of Physical Adversarial Patches,” ICCV 2025 Workshop. [CVF Open Access](https://openaccess.thecvf.com/content/ICCV2025W/Findings/html/Wang_BadPatch_Diffusion-Based_Generation_of_Physical_Adversarial_Patches_ICCVW_2025_paper.html). Nguồn tác giả nêu 1.131 ảnh. [Trang dự án](https://github.com/Wwangb/AdvT-shirt-1K).

[3] M. Lei et al., “YOLOv13: Real-Time Object Detection with Hypergraph-Enhanced Adaptive Visual Perception,” arXiv:2506.17733, 2025. [Bài báo](https://arxiv.org/abs/2506.17733).

## Phạm vi dữ liệu

Bản mirror có 1.123 ảnh và 1.548 hộp patch theo kiểm kê; nguồn tác giả nêu 1.131 ảnh. Dataset đề tài chưa có ảnh/nhãn trong các split. Mọi kết quả ở đây là kết quả proxy AdvT-shirt-1K, không phải kết quả xác nhận trên bộ dữ liệu NCKH gốc. Chi tiết nguồn và giấy phép nằm ở results/external_dataset_sources.md.
"""
report = (front + "\n".join(bm) + "\n\n" + latex_block + "\n\n" + mid + "\n".join(th)
          + "\n\n<!-- PORTRAIT -->\n\n" + trade_text + "\n\n" + "\n".join(ah)
          + "\n\n![Hình 3. Độ chính xác và latency khi lần lượt loại bỏ hoặc thay thế thành phần](adv_tshirt_patch_results/advpatch_ablation_map_latency.png)"
          + "\n\n<!-- PORTRAIT -->\n\n" + ablation_text)
MD.write_text(report, encoding="utf-8")


def font(run, size=12, bold=None, italic=None, name="Times New Roman"):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), name)
    rpr.rFonts.set(qn("w:hAnsi"), name)
    rpr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic


def rich(par, text, size=12, code=False):
    pat = re.compile(r"(\*\*.+?\*\*|\[[^\]]+\]\([^)]+\))")
    pos = 0
    for m in pat.finditer(text):
        if m.start() > pos:
            r=par.add_run(text[pos:m.start()]); font(r,size=8 if code else size,name="Consolas" if code else "Times New Roman")
        token=m.group()
        if token.startswith("**"):
            r=par.add_run(token[2:-2]); font(r,size=size,bold=True)
        else:
            z=re.match(r"\[([^\]]+)\]\(([^)]+)\)",token)
            r=par.add_run(z.group(1)+" ("+z.group(2)+")"); font(r,size=size)
        pos=m.end()
    if pos<len(text):
        r=par.add_run(text[pos:]); font(r,size=8 if code else size,name="Consolas" if code else "Times New Roman")


def cells(line):
    return [v.strip() for v in line.strip().strip("|").split("|")]


def set_shading(cell, color):
    pr=cell._tc.get_or_add_tcPr()
    shd=OxmlElement("w:shd"); shd.set(qn("w:fill"),color); pr.append(shd)


def add_table(doc, raw, landscape):
    data=[cells(x) for x in raw]
    if len(data)>1 and all(re.fullmatch(r":?-{3,}:?",v.replace(" ","")) for v in data[1]): data.pop(1)
    n=max(map(len,data))
    t=doc.add_table(rows=len(data),cols=n)
    t.style="Table Grid"; t.autofit=True; t.alignment=WD_ALIGN_PARAGRAPH.CENTER
    fs=8 if n>=10 else (8.5 if n>=7 else 9.5)
    for ri,row in enumerate(data):
        for ci in range(n):
            c=t.cell(ri,ci); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            par=c.paragraphs[0]; par.paragraph_format.space_after=Pt(0)
            par.alignment=WD_ALIGN_PARAGRAPH.CENTER if ci else WD_ALIGN_PARAGRAPH.LEFT
            par.clear(); rich(par,row[ci] if ci<len(row) else "",fs)
            for r in par.runs:
                r.font.size=Pt(fs)
                if ri==0:
                    r.bold=True; r.font.color.rgb=RGBColor(255,255,255)
            if ri==0: set_shading(c,"17365D")
            elif ri%2==0: set_shading(c,"F1F5F9")
        if ri==0:
            trpr=t.rows[0]._tr.get_or_add_trPr(); hdr=OxmlElement("w:tblHeader"); hdr.set(qn("w:val"),"true"); trpr.append(hdr)
    doc.add_paragraph()


doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(21); sec.page_height=Cm(29.7)
sec.top_margin=Cm(2); sec.bottom_margin=Cm(2); sec.left_margin=Cm(2); sec.right_margin=Cm(2)
normal=doc.styles["Normal"]
normal.font.name="Times New Roman"; normal._element.rPr.rFonts.set(qn("w:ascii"),"Times New Roman"); normal._element.rPr.rFonts.set(qn("w:hAnsi"),"Times New Roman"); normal._element.rPr.rFonts.set(qn("w:eastAsia"),"Times New Roman")
normal.font.size=Pt(12); normal.paragraph_format.line_spacing=1.15; normal.paragraph_format.space_after=Pt(6)
for name,size in (("Title",16),("Heading 1",14),("Heading 2",12)):
    st=doc.styles[name]; st.font.name="Times New Roman"; st._element.rPr.rFonts.set(qn("w:ascii"),"Times New Roman"); st._element.rPr.rFonts.set(qn("w:hAnsi"),"Times New Roman"); st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor(0,0,0)
title_ppr=doc.styles["Title"]._element.get_or_add_pPr()
title_border=title_ppr.find(qn("w:pBdr"))
if title_border is not None: title_ppr.remove(title_border)
doc.core_properties.title="Đánh giá YOLOv13 cho định vị miếng dán đối kháng vật lý trên AdvT-shirt-1K"
doc.core_properties.author="Phạm Tuấn Anh; Mai Quốc Bảo"
doc.core_properties.subject="Báo cáo nghiên cứu khoa học"
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
rr=footer.add_run(); begin=OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"),"begin"); instr=OxmlElement("w:instrText"); instr.text=" PAGE "; separate=OxmlElement("w:fldChar"); separate.set(qn("w:fldCharType"),"separate"); tx=OxmlElement("w:t"); tx.text="1"; end=OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"),"end")
for n in (begin,instr,separate,tx,end): rr._r.append(n)
font(rr,9)

lines=report.splitlines(); i=0; landscape=False
while i<len(lines):
    s=lines[i].strip()
    if not s: i+=1; continue
    if s=="<!-- LANDSCAPE -->":
        sec=doc.add_section(WD_SECTION.NEW_PAGE); sec.orientation=WD_ORIENT.LANDSCAPE; sec.page_width=Cm(29.7); sec.page_height=Cm(21); sec.top_margin=Cm(1.5); sec.bottom_margin=Cm(1.5); sec.left_margin=Cm(1.5); sec.right_margin=Cm(1.5); landscape=True; i+=1; continue
    if s=="<!-- PORTRAIT -->":
        sec=doc.add_section(WD_SECTION.NEW_PAGE); sec.orientation=WD_ORIENT.PORTRAIT; sec.page_width=Cm(21); sec.page_height=Cm(29.7); sec.top_margin=Cm(2); sec.bottom_margin=Cm(2); sec.left_margin=Cm(2); sec.right_margin=Cm(2); landscape=False; i+=1; continue
    if s.startswith("|"):
        chunk=[]
        while i<len(lines) and lines[i].strip().startswith("|"): chunk.append(lines[i].strip()); i+=1
        add_table(doc,chunk,landscape); continue
    if s.startswith("!["):
        m=re.match(r"!\[([^\]]*)\]\(([^)]+)\)",s)
        if m:
            path=(OUT/m.group(2)).resolve()
            if path.is_file() and OUT.resolve() in path.parents:
                par=doc.add_paragraph(); par.alignment=WD_ALIGN_PARAGRAPH.CENTER
                par.add_run().add_picture(str(path),width=Inches(9.8 if landscape else 6.1))
                cp=doc.add_paragraph(); cp.alignment=WD_ALIGN_PARAGRAPH.CENTER; rich(cp,m.group(1),9)
        i+=1; continue
    if s.startswith("# "):
        par=doc.add_paragraph(style="Title"); par.alignment=WD_ALIGN_PARAGRAPH.CENTER; rich(par,s[2:],16); i+=1; continue
    if s.startswith("## "):
        par=doc.add_paragraph(style="Heading 1"); rich(par,s[3:],14); i+=1; continue
    if s.startswith("### "):
        par=doc.add_paragraph(style="Heading 2"); rich(par,s[4:],12); i+=1; continue
    if s.startswith(">"):
        par=doc.add_paragraph(); par.paragraph_format.left_indent=Cm(.6); par.paragraph_format.right_indent=Cm(.6); rich(par,s.lstrip("> "),10)
        for r in par.runs:r.italic=True
        i+=1; continue
    if s.startswith("- "):
        par=doc.add_paragraph(style="List Bullet"); rich(par,s[2:],11); i+=1; continue
    if re.match(r"^\d+[.)] ",s):
        par=doc.add_paragraph(style="List Number"); rich(par,re.sub(r"^\d+[.)] ","",s),11); i+=1; continue
    if lines[i].startswith("    "):
        i+=1; continue
    par=doc.add_paragraph()
    par.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT if re.match(r"^\[\d+\]",s) else WD_ALIGN_PARAGRAPH.JUSTIFY
    rich(par,s,12); i+=1

doc.save(DOCX)
print(f"Created {MD}")
print(f"Created {DOCX}")
print(f"Benchmark rows: {len(test)}; ablation rows: {len(ABL)}")

