"""Aggregate the completed multi-seed YOLO study into Markdown and DOCX."""

from __future__ import annotations

import json
import math
import hashlib
import shutil
import statistics
import csv
import re
from datetime import datetime, timezone
from pathlib import Path

from pipeline_utils import PROJECT_ROOT, configure_project_environment, project_path, require_local_venv

configure_project_environment()

MODELS = ("yolo11n", "yolov12n", "yolov13n", "yolov13s")
SEEDS = (42, 100, 2026, 777, 999)
BUDGETS = (5, 10, 20, 30, 60, 120)
MODEL_NAMES = {"yolo11n": "YOLO11n", "yolov12n": "YOLOv12n", "yolov13n": "YOLOv13n", "yolov13s": "YOLOv13s"}
T_CRIT_95 = {1: None, 2: 12.706, 3: 4.303, 4: 3.182, 5: 2.776, 6: 2.571, 7: 2.447, 8: 2.365, 9: 2.306, 10: 2.262}
OUT = project_path("results")
EXP = project_path("results/fullscale_experiments")
REPORT_CREATED_AT = datetime.now(timezone.utc)


def read_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def stats(values: list[float]) -> dict:
    values = [float(v) for v in values if v is not None and math.isfinite(float(v))]
    n = len(values)
    if not n:
        return {"n": 0, "mean": None, "sd": None, "ci95": None}
    mean = statistics.mean(values)
    sd = statistics.stdev(values) if n > 1 else 0.0
    tcrit = T_CRIT_95.get(n)
    if tcrit is None and n > 10:
        tcrit = 1.96
    half = tcrit * sd / math.sqrt(n) if tcrit is not None else None
    return {"n": n, "mean": mean, "sd": sd, "ci95": [mean - half, mean + half] if half is not None else None}


def fmt(value: float | None, digits: int = 2) -> str:
    return "—" if value is None else f"{value:.{digits}f}"


def fmt_stats(summary: dict, scale: float = 1.0, digits: int = 2, unit: str = "") -> str:
    if summary["n"] == 0:
        return "—"
    mean, sd = summary["mean"] * scale, summary["sd"] * scale
    ci = summary["ci95"]
    if ci is None:
        ci_text = "95% CI n/a"
    else:
        ci_text = f"95% CI [{ci[0] * scale:.{digits}f}, {ci[1] * scale:.{digits}f}]"
    suffix = f" {unit}" if unit else ""
    return f"{mean:.{digits}f} ± {sd:.{digits}f}{suffix}; {ci_text}"


def fmt_latex_stats(summary: dict, scale: float = 1.0, digits: int = 1) -> str:
    if not summary["n"]:
        return "--"
    mean, sd = summary["mean"] * scale, summary["sd"] * scale
    ci = summary["ci95"]
    interval = f"CI [{ci[0] * scale:.{digits}f}, {ci[1] * scale:.{digits}f}]" if ci else "CI n/a"
    return r"\shortstack{" + f"${mean:.{digits}f} \\pm {sd:.{digits}f}$" + r"\\" + interval + "}"


def load_records() -> dict[str, dict[int, dict]]:
    records = {}
    missing = []
    for model in MODELS:
        records[model] = {}
        for seed in SEEDS:
            path = EXP / "records" / model / f"seed-{seed}.json"
            if not path.is_file():
                missing.append(str(path.relative_to(PROJECT_ROOT)))
                continue
            record = read_json(path)
            if record.get("status") != "complete":
                raise RuntimeError(f"Training record is not complete: {path}")
            records[model][seed] = record
    if missing:
        raise RuntimeError("Training records are missing: " + ", ".join(missing))
    return records


def archive_epoch_histories(records: dict) -> None:
    """Keep compact per-epoch metric CSVs with the aggregated experiment evidence."""
    for model in MODELS:
        for seed in SEEDS:
            record = records[model][seed]
            # Earlier runner versions sampled this timer before test evaluation.
            # Retain the measured value under a name describing its actual scope.
            if "process_wall_seconds_including_model_load_and_test" in record:
                record["process_wall_seconds_before_snapshot_test"] = record.pop("process_wall_seconds_including_model_load_and_test")
                record["wall_timer_scope_correction"] = "The original field name included test, but its timer was sampled before snapshot test evaluation."
            source = project_path(record["results_csv"], must_exist=True)
            target = EXP / "epoch_history" / model / f"seed-{seed}.csv"
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(source, target)
            with source.open("r", encoding="utf-8-sig", newline="") as handle:
                history = list(csv.DictReader(handle))
            by_epoch = {int(float(row["epoch"])): row for row in history if row.get("epoch")}
            for milestone in record.get("milestones", []):
                endpoint_row = by_epoch.get(int(milestone["endpoint_epoch"])) if milestone.get("endpoint_epoch") is not None else None
                if endpoint_row is None:
                    milestone["endpoint_validation"] = None
                    continue
                milestone["endpoint_validation"] = {
                    "precision": float(endpoint_row["metrics/precision(B)"]),
                    "recall": float(endpoint_row["metrics/recall(B)"]),
                    "map50": float(endpoint_row["metrics/mAP50(B)"]),
                    "map50_95": float(endpoint_row["metrics/mAP50-95(B)"]),
                }
            digest = hashlib.sha256(target.read_bytes()).hexdigest()
            record["epoch_history_csv"] = str(target.relative_to(PROJECT_ROOT)).replace("\\", "/")
            record["epoch_history_sha256"] = digest
            record_path = EXP / "records" / model / f"seed-{seed}.json"
            record_path.write_text(json.dumps(record, ensure_ascii=False, indent=2), encoding="utf-8")


def milestone_for(record: dict, budget: int) -> dict | None:
    return next((r for r in record.get("milestones", []) if r.get("budget_minutes") == budget), None)


def validation_curve(records: dict) -> dict[str, dict[int, dict]]:
    curves = {}
    for model in MODELS:
        curves[model] = {}
        for budget in BUDGETS:
            rows = [milestone_for(records[model][seed], budget) for seed in SEEDS]
            rows = [r for r in rows if r and r.get("status") == "complete" and r.get("validation", {}).get("map50_95") is not None]
            curves[model][budget] = {
                "map": stats([r["validation"]["map50_95"] for r in rows]),
                "endpoint_map": stats([r["endpoint_validation"]["map50_95"] for r in rows if r.get("endpoint_validation")]),
                "rows": rows,
                "train_loss": stats([sum(r["endpoint_loss"][k] or 0 for k in ("train_box", "train_cls", "train_dfl")) for r in rows if r.get("endpoint_loss")]),
                "val_loss": stats([sum(r["endpoint_loss"][k] or 0 for k in ("val_box", "val_cls", "val_dfl")) for r in rows if r.get("endpoint_loss")]),
            }
    return curves


def select_sweet_spots(curves: dict) -> dict[str, dict]:
    sweet = {}
    for model in MODELS:
        means = {b: curves[model][b]["map"]["mean"] for b in BUDGETS if curves[model][b]["map"]["n"] == len(SEEDS)}
        if not means:
            raise RuntimeError(f"No milestone has a complete 5-seed validation mean for {model}")
        peak_budget = max(means, key=means.get)
        peak_mean = means[peak_budget]
        budget = next(b for b in BUDGETS if b in means and peak_mean - means[b] <= 0.005)
        sweet[model] = {
            "sweet_budget": budget,
            "sweet_mean": means[budget],
            "sweet_stats": curves[model][budget]["map"],
            "peak_budget": peak_budget,
            "peak_mean": peak_mean,
            "extra_gain_pp": (peak_mean - means[budget]) * 100,
            "means": means,
        }
    return sweet


def fit_diagnostics(records: dict) -> dict[str, dict]:
    diagnostics = {}
    for model in MODELS:
        train_deltas, val_deltas, late_map_gains = [], [], []
        overfit_candidates = 0
        continued_learning = 0
        paired = 0
        for seed in SEEDS:
            early = milestone_for(records[model][seed], 5)
            mid = milestone_for(records[model][seed], 60)
            late = milestone_for(records[model][seed], 120)
            if not all(row and row.get("status") == "complete" for row in (early, mid, late)):
                continue
            paired += 1
            early_loss, late_loss = early.get("endpoint_loss") or {}, late.get("endpoint_loss") or {}
            train_early = sum(early_loss.get(k) or 0 for k in ("train_box", "train_cls", "train_dfl"))
            train_late = sum(late_loss.get(k) or 0 for k in ("train_box", "train_cls", "train_dfl"))
            val_early = sum(early_loss.get(k) or 0 for k in ("val_box", "val_cls", "val_dfl"))
            val_late = sum(late_loss.get(k) or 0 for k in ("val_box", "val_cls", "val_dfl"))
            train_deltas.append(train_late - train_early)
            val_deltas.append(val_late - val_early)
            early_map = (early.get("endpoint_validation") or {}).get("map50_95")
            mid_map = (mid.get("endpoint_validation") or {}).get("map50_95")
            late_map = (late.get("endpoint_validation") or {}).get("map50_95")
            if None not in (early_map, mid_map, late_map):
                gain = (late_map - mid_map) * 100
                late_map_gains.append(gain)
                if gain > 0.5:
                    continued_learning += 1
                if train_late < train_early * 0.95 and val_late > val_early * 1.05 and (late_map - early_map) <= 0.005:
                    overfit_candidates += 1
        diagnostics[model] = {
            "paired_seeds": paired,
            "mean_train_loss_change_5_to_120": stats(train_deltas),
            "mean_val_loss_change_5_to_120": stats(val_deltas),
            "mean_val_map_gain_60_to_120_pp": stats(late_map_gains),
            "continued_learning_seeds": continued_learning,
            "overfit_candidate_seeds": overfit_candidates,
        }
    return diagnostics


def load_latency() -> dict[str, dict[int, dict]]:
    path = EXP / "latency_records.json"
    if not path.is_file():
        raise RuntimeError("Latency records are missing; run benchmark_fullscale_latency.py --all after training.")
    data = read_json(path)
    result = {m: {} for m in MODELS}
    for row in data.get("records", []):
        if row.get("timing_precision") != "fp32" or row.get("batch") != 1 or row.get("imgsz") != 416:
            raise RuntimeError("Latency records must use the declared FP32 batch-1 416px setup")
        if row.get("warmup_frames") != 32 or row.get("measured_frames") != 500:
            raise RuntimeError("Latency records must contain 32 warmup and 500 measured frames")
        value = row.get("latency_median_ms_per_frame")
        fps = row.get("fps_from_median_latency")
        if value is None or fps is None or not math.isfinite(value) or value <= 0 or not math.isfinite(fps) or not math.isclose(fps, 1000 / value):
            raise RuntimeError("Invalid latency/FPS pair in benchmark records")
        result[row["model_key"]][int(row["seed"])] = row
    for model in MODELS:
        if set(result[model]) != set(SEEDS):
            raise RuntimeError(f"Latency records for {model} do not include all five seeds")
    return result


def metric_values(records: dict, model: str, budget: int, split: str, metric: str) -> list[float]:
    values = []
    for seed in SEEDS:
        row = milestone_for(records[model][seed], budget)
        if row and row.get("status") == "complete" and row.get(split) and row[split].get(metric) is not None:
            values.append(row[split][metric])
    return values


def generate_plot(curves: dict) -> Path:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    path = OUT / "fullscale_mAP_growth.png"
    fig, axes = plt.subplots(1, 2, figsize=(11, 5.2), sharey=True, constrained_layout=True)
    colors = {"yolo11n": "#2563eb", "yolov12n": "#ea580c", "yolov13n": "#16a34a", "yolov13s": "#9333ea"}
    panels = ((axes[0], "map", "Best validation checkpoint"), (axes[1], "endpoint_map", "Last completed epoch"))
    for ax, metric_key, title in panels:
        for model in MODELS:
            x, y, low, high = [], [], [], []
            for budget in BUDGETS:
                summary = curves[model][budget][metric_key]
                if not summary["n"]:
                    continue
                x.append(budget)
                y.append(summary["mean"] * 100)
                ci = summary["ci95"]
                low.append((summary["mean"] - ci[0]) * 100 if ci else 0)
                high.append((ci[1] - summary["mean"]) * 100 if ci else 0)
                if summary["n"] < len(SEEDS):
                    ax.annotate(f"n={summary['n']}", (budget, summary["mean"] * 100), xytext=(4, 5), textcoords="offset points", fontsize=7, color=colors[model])
            ax.errorbar(x, y, yerr=[low, high], marker="o", linewidth=1.6, capsize=2, label=MODEL_NAMES[model], color=colors[model])
        ax.set_xscale("log", base=2)
        ax.set_xticks(BUDGETS, labels=[str(b) for b in BUDGETS])
        ax.set_xlabel("Training budget (minutes; log scale)")
        ax.set_title(title)
        ax.grid(True, which="both", alpha=0.25)
        ax.legend(frameon=False, fontsize=8)
    axes[0].set_ylabel("Validation mAP@0.5:0.95 (%)")
    fig.suptitle("Validation performance by budget: selected checkpoint vs. training endpoint")
    fig.savefig(path, dpi=180)
    plt.close(fig)
    return path


def markdown_table1(records: dict, sweet: dict, latency: dict) -> tuple[str, str]:
    headers = ["Model", "Params (M)", "FLOPs (G)", "Precision (%)", "Recall (%)", "mAP@0.5 (%)", "mAP@0.5:0.95 (%)", "Latency (ms/frame)", "FPS"]
    lines = ["| " + " | ".join(headers) + " |", "|" + "|".join(["---"] * len(headers)) + "|"]
    latex_rows = []
    for model in MODELS:
        budget = sweet[model]["sweet_budget"]
        rows = [milestone_for(records[model][s], budget) for s in SEEDS]
        metrics = {
            key: stats([r["test"][key] for r in rows if r and r.get("test") and r["test"].get(key) is not None])
            for key in ("precision", "recall", "map50", "map50_95")
        }
        latency_stats = stats([latency[model][s]["latency_median_ms_per_frame"] for s in SEEDS])
        fps_stats = stats([latency[model][s]["fps_from_median_latency"] for s in SEEDS])
        params = stats([r["parameters_m"] for r in rows if r and r.get("parameters_m") is not None])
        flops = stats([r["flops_g"] for r in rows if r and r.get("flops_g") is not None])
        cells = [
            MODEL_NAMES[model], fmt_stats(params, 1, 3), fmt_stats(flops, 1, 2),
            fmt_stats(metrics["precision"], 100, 2), fmt_stats(metrics["recall"], 100, 2),
            fmt_stats(metrics["map50"], 100, 2), fmt_stats(metrics["map50_95"], 100, 2),
            fmt_stats(latency_stats, 1, 2, "ms"), fmt_stats(fps_stats, 1, 2),
        ]
        lines.append("| " + " | ".join(cells) + " |")
        compact = [
            MODEL_NAMES[model], fmt(params["mean"], 3), fmt(flops["mean"], 2),
            fmt_latex_stats(metrics["precision"], 100, 1), fmt_latex_stats(metrics["recall"], 100, 1),
            fmt_latex_stats(metrics["map50"], 100, 1), fmt_latex_stats(metrics["map50_95"], 100, 1),
            fmt_latex_stats(latency_stats, 1, 1), fmt_latex_stats(fps_stats, 1, 1),
        ]
        latex_rows.append(" & ".join(v.replace("±", r"$\pm$").replace("; 95% CI ", "; CI ") for v in compact) + r" \\")
    latex = """% IEEE-style comparison table (compile in a document using IEEEtran)
% Preamble dependency: booktabs package for top/mid/bottom rules.
\\begin{table*}[t]
\\centering
\\caption{Multi-seed test performance and batch-1 latency at the validation-selected sweet spot. Values are mean $\\pm$ SD with 95\\% t confidence intervals.}
\\label{tab:sweetspot}
\\scriptsize
\\setlength{\\tabcolsep}{2pt}
\\begin{tabular}{lrrrrrrrr}
\\toprule
Model & Params (M) & FLOPs (G) & Precision (\\%) & Recall (\\%) & mAP@0.5 (\\%) & mAP@0.5:0.95 (\\%) & Latency (ms) & FPS \\\\
\\midrule
""" + "\n".join(latex_rows) + "\n" + r"""\bottomrule
\end{tabular}
\end{table*}"""
    return "\n".join(lines), latex


def markdown_table2(records: dict, curves: dict, sweet: dict) -> tuple[str, str]:
    headers = ["Model", "Budget (min)", "n", "Best-checkpoint val mAP95 (%) μ±σ; 95% CI", "Last-epoch val mAP95 (%) μ±σ; 95% CI", "Δ best from previous (pp)", "Best-checkpoint gain (pp/min)", "Train loss (sum)", "Val loss (sum)"]
    lines = ["| " + " | ".join(headers) + " |", "|" + "|".join(["---"] * len(headers)) + "|"]
    latex_rows = []
    for model in MODELS:
        previous_budget = None
        previous_mean = None
        previous_n = None
        for budget in BUDGETS:
            cell = curves[model][budget]
            s = cell["map"]
            chosen = sweet[model]["sweet_budget"] == budget
            display_model = f"**{MODEL_NAMES[model]}**" if chosen else MODEL_NAMES[model]
            if not s["n"]:
                lines.append(f"| {display_model} | {budget} | 0 | — | — | — | — | — | — |")
                latex_rows.append(" & ".join([MODEL_NAMES[model], str(budget), "0", "--", "--", "--", "--", "--", "--"]) + r" \\")
                continue
            delta = (s["mean"] - previous_mean) * 100 if previous_mean is not None and s["n"] == previous_n == len(SEEDS) else None
            dt = budget - previous_budget if previous_budget is not None else None
            gain = delta / dt if delta is not None and dt else None
            train_loss = cell["train_loss"]
            val_loss = cell["val_loss"]
            map_text = fmt_stats(s, 100, 2)
            endpoint_text = fmt_stats(cell["endpoint_map"], 100, 2)
            lines.append("| {} | {}{} | {} | {} | {} | {} | {} | {} | {} |".format(
                display_model, f"**{budget}**" if chosen else budget, "", s["n"], map_text, endpoint_text,
                fmt(delta), fmt(gain), fmt(train_loss["mean"]), fmt(val_loss["mean"])))
            latex_rows.append(" & ".join([
                MODEL_NAMES[model], str(budget) + (r"$^\star$" if chosen else ""), str(s["n"]),
                fmt_latex_stats(s, 100, 1),
                fmt_latex_stats(cell["endpoint_map"], 100, 1),
                fmt(delta), fmt(gain), fmt(train_loss["mean"]), fmt(val_loss["mean"]),
            ]) + r" \\")
            previous_budget, previous_mean = budget, s["mean"]
            previous_n = s["n"]
    latex = """% IEEE-style convergence table; a star marks the earliest budget within 0.5 pp of peak mean validation mAP.
% Preamble dependency: booktabs package for top/mid/bottom rules.
\\begin{table*}[t]
\\centering
\\caption{Validation convergence and endpoint loss at each training budget.}
\\label{tab:convergence}
\\scriptsize
\\begin{tabular}{llrrrrrrr}
\\toprule
Model & Minutes & $n$ & Best-checkpoint val mAP95 (\\%) & Last-epoch val mAP95 (\\%) & $\\Delta$ best (pp) & Gain (pp/min) & Train loss & Val loss \\\\
\\midrule
""" + "\n".join(latex_rows) + "\n" + r"""\bottomrule
\end{tabular}
\end{table*}"""
    return "\n".join(lines), latex


def markdown_fit_table(diagnostics: dict) -> str:
    lines = [
        "| Model | Paired seeds | Δ train loss (5→120m) | Δ val loss (5→120m) | Δ endpoint val mAP95 (60→120m, pp) | Seeds still gaining >0.5 pp | Overfit candidate seeds |",
        "|---|---:|---:|---:|---:|---:|---:|",
    ]
    for model in MODELS:
        d = diagnostics[model]
        lines.append("| {} | {} | {} | {} | {} | {} | {} |".format(
            MODEL_NAMES[model], d["paired_seeds"],
            fmt(d["mean_train_loss_change_5_to_120"]["mean"]),
            fmt(d["mean_val_loss_change_5_to_120"]["mean"]),
            fmt(d["mean_val_map_gain_60_to_120_pp"]["mean"]),
            d["continued_learning_seeds"], d["overfit_candidate_seeds"],
        ))
    return "\n".join(lines)


def make_report(records: dict) -> tuple[str, dict, dict, dict, dict]:
    curves = validation_curve(records)
    sweet = select_sweet_spots(curves)
    diagnostics = fit_diagnostics(records)
    latency = load_latency()
    for model in MODELS:
        for seed in SEEDS:
            row = latency[model][seed]
            selected = milestone_for(records[model][seed], sweet[model]["sweet_budget"])
            if row.get("sweet_spot_minutes") != sweet[model]["sweet_budget"] or row.get("checkpoint") != selected.get("checkpoint"):
                raise RuntimeError(f"Latency checkpoint does not match the selected sweet spot: {model}, seed {seed}")
    plot_path = generate_plot(curves)
    table1, latex1 = markdown_table1(records, sweet, latency)
    table2, latex2 = markdown_table2(records, curves, sweet)
    fit_table = markdown_fit_table(diagnostics)

    dataset_summary = read_json(project_path("results/fullscale_setup/dataset_build_summary.json", must_exist=True))
    audit = read_json(project_path("results/fullscale_setup/fullscale_dataset_audit.json", must_exist=True))
    integrity = read_json(project_path("results/fullscale_setup/split_integrity_audit.json", must_exist=True))
    dataset_manifest = read_json(project_path("dataset/fullscale_person/manifest.json", must_exist=True))
    lat_means = {m: stats([latency[m][s]["latency_median_ms_per_frame"] for s in SEEDS]) for m in MODELS}
    best_test_model = max(MODELS, key=lambda m: stats(metric_values(records, m, sweet[m]["sweet_budget"], "test", "map50_95"))["mean"])
    peak_gains = [sweet[m]["extra_gain_pp"] for m in MODELS]
    plateau_models = [MODEL_NAMES[m] for m in MODELS if sweet[m]["extra_gain_pp"] <= 0.5 + 1e-9]

    hardware = records[MODELS[0]][SEEDS[0]]["hardware_software"]
    try:
        machine = read_json(project_path("results/fullscale_setup/hardware.json"))
    except FileNotFoundError:
        machine = {}
    split_counts = dataset_summary["split_counts"]
    total_images = dataset_summary["total_images"]
    audit_invalid = sum(v["invalid_label_rows"] for v in audit["splits"].values())
    sanitization = dataset_manifest.get("advtshirt_box_sanitization", {})
    clipped_rows = sanitization.get("source_out_of_bounds_rows", 0)
    reencoded_rows = sanitization.get("rows_reencoded_or_edge_adjusted", 0)
    max_overrun = sanitization.get("maximum_source_overrun_normalized", 0.0)
    sweet_lines = []
    for model in MODELS:
        item = sweet[model]
        ci = item["sweet_stats"]["ci95"]
        sweet_lines.append(
            f"- **{MODEL_NAMES[model]}:** sweet spot {item['sweet_budget']} phút; peak quan sát ở {item['peak_budget']} phút; "
            f"validation mAP@0.5:0.95 trung bình {item['sweet_mean'] * 100:.2f}% "
            f"(95% CI [{ci[0] * 100:.2f}, {ci[1] * 100:.2f}]%), cách peak {item['extra_gain_pp']:.2f} điểm phần trăm."
        )
    best_test = stats(metric_values(records, best_test_model, sweet[best_test_model]["sweet_budget"], "test", "map50_95"))
    illustrative_deadline_ms = 1000 / 30
    realtime_candidates = [m for m in MODELS if lat_means[m]["mean"] <= illustrative_deadline_ms]
    if realtime_candidates:
        realtime_choice = max(realtime_candidates, key=lambda m: stats(metric_values(records, m, sweet[m]["sweet_budget"], "test", "map50_95"))["mean"])
        realtime_test = stats(metric_values(records, realtime_choice, sweet[realtime_choice]["sweet_budget"], "test", "map50_95"))
        realtime_text = (f"Nếu lấy mục tiêu 30 FPS làm ví dụ minh họa (ngân sách pipeline {illustrative_deadline_ms:.2f} ms/frame), "
                         f"{MODEL_NAMES[realtime_choice]} có test mAP95 trung bình cao nhất trong các model có latency trung bình dưới ngưỡng: "
                         f"{realtime_test['mean'] * 100:.2f}% tại {sweet[realtime_choice]['sweet_budget']} phút, latency {lat_means[realtime_choice]['mean']:.2f} ms/frame. "
                         "Đây là sàng lọc theo trung bình trên RTX 5060, chưa bảo đảm deadline từng frame hoặc 30 FPS của toàn ứng dụng; cần đo tail latency, decode và thiết bị đích.")
    else:
        realtime_text = "Không model nào có latency pipeline trung bình dưới ngưỡng minh họa 33.33 ms/frame cho 30 FPS trong cấu hình này."
    accuracy_choice_text = (f"Khi ưu tiên độ chính xác trong nhóm đã đo, chọn {MODEL_NAMES[best_test_model]} với ngân sách {sweet[best_test_model]['sweet_budget']} phút: "
                            f"test mAP95 {best_test['mean'] * 100:.2f}%, latency pipeline trung bình {lat_means[best_test_model]['mean']:.2f} ms/frame ở FP32 batch 1. "
                            "Đối chiếu mức latency này với yêu cầu tốc độ trước khi triển khai; định dạng FP16, ONNX hoặc TensorRT chưa được benchmark trong nghiên cứu này.")
    extra_time_text = "; ".join(f"{MODEL_NAMES[m]} +{sweet[m]['extra_gain_pp']:.2f} pp" for m in MODELS)
    cpu = machine.get("cpu_name", "chưa ghi nhận")
    ram = machine.get("ram_gib", "chưa ghi nhận")
    split_ratio = ", ".join(f"{k} {v['images']} ảnh" for k, v in split_counts.items())
    actual_training_minutes = [records[m][s]["actual_training_seconds"] / 60 for m in MODELS for s in SEEDS]
    total_training_hours = sum(actual_training_minutes) / 60
    max_cap_overrun_minutes = max(max(0, duration - max(BUDGETS)) for duration in actual_training_minutes)
    pretrain_source_test = integrity["source_split_counts"]["test"].get("COCO 2017:train2017", 0)
    duplicate_count = len(integrity["cross_split_exact_duplicate_groups"])
    early_selection_lines = []
    for model in MODELS:
        selected_rows = [milestone_for(records[model][seed], sweet[model]["sweet_budget"]) for seed in SEEDS]
        first_count = sum(row["selected_epoch"] == 1 for row in selected_rows)
        selected_times = [row["selected_elapsed_seconds"] / 60 for row in selected_rows]
        early_selection_lines.append(f"{MODEL_NAMES[model]} chọn epoch 1 ở {first_count}/5 seed tại sweet spot; thời gian checkpoint thực được chọn nằm trong {min(selected_times):.2f}–{max(selected_times):.2f} phút")
    selection_text = "; ".join(early_selection_lines) + "."
    stability_text = "; ".join(
        f"{MODEL_NAMES[m]} có SD test mAP95 {stats(metric_values(records, m, sweet[m]['sweet_budget'], 'test', 'map50_95'))['sd'] * 100:.2f} điểm phần trăm"
        for m in MODELS
    ) + "."
    comparison_lines = []
    for candidate in ("yolov13n", "yolov13s"):
        for baseline in ("yolo11n", "yolov12n"):
            differences = [
                (milestone_for(records[candidate][seed], sweet[candidate]["sweet_budget"])["test"]["map50_95"]
                 - milestone_for(records[baseline][seed], sweet[baseline]["sweet_budget"])["test"]["map50_95"]) * 100
                for seed in SEEDS
            ]
            difference = stats(differences)
            comparison_lines.append(f"- {MODEL_NAMES[candidate]} − {MODEL_NAMES[baseline]} tại sweet spot riêng: Δ test mAP95 {fmt_stats(difference)} điểm phần trăm; latency trung bình {lat_means[candidate]['mean']:.2f} so với {lat_means[baseline]['mean']:.2f} ms/frame.")
    boundary_models = [MODEL_NAMES[m] for m in MODELS if sweet[m]["sweet_budget"] == max(BUDGETS)]
    boundary_text = (
        f"{', '.join(boundary_models)} có sweet spot tại biên 120 phút; kết quả đó chưa chứng minh bão hòa và cần chạy dài hơn để tìm trần ngoài cửa sổ quan sát."
        if boundary_models else "Không có sweet spot nằm ở biên 120 phút; kết luận vẫn giới hạn trong cửa sổ quan sát."
    )
    markdown = f"""# Báo cáo thực nghiệm quy mô lớn: hội tụ thời gian và hiệu quả triển khai YOLOv13

**Ngày lập:** {datetime.now().astimezone().strftime('%d-%m-%Y')}  
**Thiết bị huấn luyện:** {hardware.get('gpu', 'NVIDIA GeForce RTX 5060')} ({hardware.get('gpu_memory_gib', '8.15')} GiB VRAM); **CPU:** {cpu}; **RAM:** {ram} GiB  
**Môi trường:** `.venv`; Python {hardware.get('python', '').split()[0]}; PyTorch {hardware.get('torch', 'n/a')}; Ultralytics {hardware.get('ultralytics', 'n/a')}  
**Thiết kế:** 4 mô hình × 5 seed × tối đa 120 phút; các mốc 5, 10, 20, 30, 60 và 120 phút.

## Tóm tắt

Nghiên cứu đánh giá YOLO11n, YOLOv12n, YOLOv13n và YOLOv13s trên bài toán phát hiện một lớp `person`, dùng 10.000 ảnh COCO 2017 được lấy mẫu xác định cùng toàn bộ 1.123 ảnh AdvT-shirt-1K có trong bản dữ liệu cục bộ. Cùng một split 70/15/15 được giữ cố định; năm seed độc lập được dùng cho mỗi mô hình. Tại từng mốc, checkpoint được chọn chỉ bằng validation mAP@0.5:0.95 rồi đo một lần trên test. Sweet spot được định nghĩa trước là mốc sớm nhất có trung bình validation mAP không cách peak quan sát quá 0.5 điểm phần trăm.

Theo tiêu chí đó, các sweet spot quan sát được là:

{chr(10).join(sweet_lines)}

Ở sweet spot theo validation, **{MODEL_NAMES[best_test_model]}** đạt trung bình test mAP@0.5:0.95 cao nhất trong nhóm này ({best_test['mean'] * 100:.2f}%). Kết luận chỉ áp dụng cho tập một lớp, cách lấy mẫu, cấu hình và GPU nêu trong báo cáo; nó không khẳng định ưu thế phổ quát của một phiên bản YOLO.

## 1. Thiết lập thực nghiệm

### Dữ liệu và nhãn

- COCO 2017 train/val ([nguồn chính thức](https://cocodataset.org/dataset/detection-2017.htm)) được gộp nguồn rồi lấy mẫu xác định gồm 8.000 ảnh có ít nhất một hộp người không crowd và 2.000 ảnh không có bất kỳ chú thích người nào; ảnh crowd-only bị loại khỏi hai nhóm. Chỉ category `person` được giữ.
- Bản AdvT-shirt-1K ([trang dự án](https://github.com/Wwangb/AdvT-shirt-1K), [dataset mirror](https://huggingface.co/datasets/xingjunm/AdvT-shirt)) cục bộ có 1.123 ảnh, ít hơn 8 ảnh so với 1.131 ảnh nêu trên trang dự án. Hộp người được ánh xạ sang class `person`; nhãn patch vẫn được giữ ở thư mục nguồn và không gộp vào target person.
- {clipped_rows} hộp người AdvT-shirt trong nhãn nguồn vượt biên tối đa {max_overrun:.8g} tọa độ chuẩn hóa; {reencoded_rows} dòng được tái mã hóa với đệm 1e-5 để bảo đảm hộp nằm trong ảnh. Kiểm toán cuối cùng ghi nhận {audit_invalid} dòng nhãn lỗi; tệp thiếu/thừa: 0.
- Tổng cộng {total_images:,} ảnh; chia cố định theo seed 2026: {split_ratio}. Dữ liệu dẫn xuất và ảnh COCO chỉ được giữ cục bộ trong `project_yolov13/`.
- Kiểm tra SHA-256 trên {integrity['images_hashed']:,} ảnh phát hiện {duplicate_count} nhóm ảnh trùng byte giữa các split; số nhóm trùng COCO image ID giữa split là {len(integrity['cross_split_coco_identity_groups'])}. Kiểm tra này không loại trừ khung hình gần giống hoặc cùng người/cảnh xuất hiện ở nhiều split.
- COCO annotations được phân phối riêng với license ảnh Flickr; do đó ảnh và dữ liệu thô không được đưa lên GitHub. Bản mirror AdvT-shirt khai báo Apache-2.0 trên dataset card; cần kiểm tra điều khoản nguồn trước khi tái phân phối.

### Mô hình và điều kiện huấn luyện

Mỗi seed bắt đầu từ trọng số pretrained tương ứng: [YOLO11](https://docs.ultralytics.com/models/yolo11/), [YOLOv12](https://github.com/sunsmarterjie/yolov12), [YOLOv13](https://github.com/iMoonLab/yolov13). Dùng ảnh 416 px, batch 8, SGD (lr0=0.01, lrf=0.01, momentum=0.937, weight decay=0.0005), AMP, deterministic mode, 0 dataloader workers và cùng augmentation cho mọi mô hình (mosaic=1, HSV 0.015/0.7/0.4, translate=0.1, scale=0.5, fliplr=0.5; mixup=0). Huấn luyện tiếp tục theo epoch cho đến khi vượt ngân sách 120 phút rồi dừng ở ranh giới epoch đã lưu. Mỗi epoch có checkpoint. Mốc thời gian lấy từ thời gian huấn luyện tích lũy trong `results.csv`; ở mỗi mốc, checkpoint có validation mAP@0.5:0.95 cao nhất trong các epoch đã hoàn tất được chọn. Không có epoch nào trước ngân sách thì mốc được đánh dấu thiếu dữ liệu.

Tổng thời gian tích lũy trong CSV của 20 quỹ đạo là {total_training_hours:.2f} giờ; mỗi lượt thực chạy {min(actual_training_minutes):.2f}–{max(actual_training_minutes):.2f} phút, vượt cap tối đa {max_cap_overrun_minutes:.2f} phút do ranh giới epoch. Timer này bao gồm train, validation và lưu checkpoint trong vòng lặp epoch; các phép đánh giá test sau huấn luyện và benchmark latency được thực hiện riêng. Epoch vượt ngân sách không được dùng cho mốc 120 phút.

Các mô hình YOLOv12 và YOLOv13 được nạp từ mã nguồn tương ứng đã lưu trong thư mục dự án; YOLO11 dùng Ultralytics cục bộ. Máy: GPU {hardware.get('gpu', 'NVIDIA GeForce RTX 5060')} ({hardware.get('gpu_memory_gib', 8.15)} GiB), CPU {cpu}, RAM {ram} GiB; Python {hardware.get('python', '').split()[0]}, PyTorch {hardware.get('torch', 'n/a')}, CUDA có sẵn. Giá trị CPU/RAM dựa trên máy chạy thí nghiệm.

Log của YOLOv12/YOLOv13 ghi nhận không có gói `flash_attn` và chuyển sang `torch.nn.functional.scaled_dot_product_attention`. Benchmark phản ánh implementation SDPA trong môi trường này; không dùng tốc độ công bố với một backend khác để thay thế kết quả thực đo.

CSV theo epoch của từng seed được lưu cùng checksum SHA-256 dưới `results/fullscale_experiments/epoch_history/`; checkpoint đầy đủ được giữ trong thư mục `runs/train/` của dự án.

Danh sách phiên bản dependency, commit nguồn nếu có, SHA-256 của 324 tệp Python thuộc hai implementation, bốn trọng số pretrained và manifest dữ liệu được lưu trong `results/fullscale_setup/dependency_versions.txt` và `provenance.json`. Các dấu vết này định danh môi trường đã chạy; nghiên cứu chưa thực hiện một lượt tái lập độc lập hoàn chỉnh.

### Chỉ số và thống kê

Precision, recall, mAP@0.5 và mAP@0.5:0.95 đo trên test tại checkpoint validation-selected. Test được giữ khỏi fine-tuning và chọn checkpoint; tính độc lập với pretraining không được bảo đảm.

Latency đo riêng với batch 1, 416 px, FP32 (`half=False`), bỏ 32 frame warm-up rồi lấy median thời gian tiền xử lý + inference + hậu xử lý do Ultralytics ghi nhận trên tối đa 500 ảnh test cố định; FPS = 1000 / latency median. Đây là tốc độ pipeline xử lý ảnh, không bao gồm đọc/giải mã ảnh, camera, truyền mạng hoặc toàn bộ độ trễ ứng dụng.

Parameters là tổng `numel` của tham số checkpoint chia 10⁶. FLOPs là ước tính từ `get_flops`: THOP MACs nhân 2, thường profiling ở kích thước stride rồi quy đổi theo diện tích đến 416 px; nếu nhánh đó lỗi, hàm thử trực tiếp ở 416 px. Quy tắc đếm của profiler và phép quy đổi có thể không bao phủ mọi phép toán functional/custom hoặc độ phức tạp theo kích thước của attention. Vì vậy FLOPs được dùng như ước tính tham khảo; tốc độ triển khai dựa trên latency thực đo, và không so trực tiếp với FLOPs công bố ở 640 px.

Các giá trị đa seed trình bày μ ± σ mẫu; khoảng tin cậy 95% dùng phân phối t Student với df=n−1 (n=5, df=4 tại sweet spot). Mốc có seed chưa hoàn tất epoch dùng số seed thực có trong bảng và không được chọn làm sweet spot nếu n<5. CI chưa hiệu chỉnh cho việc chọn mốc trên cùng validation. Dung sai 0.5 điểm phần trăm được áp dụng cho trung bình điểm, không phải kiểm định tương đương thống kê.

## 2. Bảng so sánh đa seed tại sweet spot

Mốc sweet spot được chọn theo validation; các chỉ số kiểm tra bên dưới dùng test holdout đối với fine-tuning. Ô metric ghi μ ± σ và 95% CI của trung bình.

{table1}

## 3. Hội tụ theo thời gian và điểm tối ưu

Bảng và Hình 1 tách hai câu hỏi: checkpoint validation tốt nhất đã có thể chọn trong ngân sách (best-so-far) và hiệu năng validation của epoch cuối cùng thực sự hoàn tất trước ngân sách. Best-so-far có thể lặp lại giữa nhiều mốc, nhất là khi trọng số khởi tạo đã mạnh; đường epoch cuối cho thấy tác động của quá trình fine-tuning. Δ và lợi ích biên trong bảng tính theo best-so-far chỉ khi cả hai mốc liền trước và hiện tại có đủ cùng năm seed; ô còn lại ghi thiếu dữ liệu để tránh so sánh các nhóm seed khác nhau. Loss lấy tại epoch cuối. Mỗi epoch có thời gian riêng nên mốc là giới hạn tối đa, không nhất thiết khớp đúng thời lượng.

{table2}

![Đường cong mAP validation trung bình đa seed và CI 95%]({plot_path.name})

### Ma trận sweet spot và bão hòa

{chr(10).join(sweet_lines)}

Theo định nghĩa dung sai, sweet spot là điểm vận hành sớm nhất gần trần validation trong cửa sổ 120 phút. Peak trung bình có thể xuất hiện muộn hơn; phần cải thiện bổ sung sau sweet spot lần lượt là {extra_time_text}. Các metric test ở từng mốc được lưu trong bản ghi JSON của seed; bảng hội tụ trình bày validation và không dùng test để chọn thời gian. Đây là sweet spot của việc chọn checkpoint trong ngân sách, không tự chứng minh rằng epoch cuối đã hội tụ.

## 4. Thảo luận

### Tăng trưởng và diminishing returns

Xem đường cong trên để đối chiếu tốc độ tăng mAP theo thời gian. Các điểm là trung bình các seed có epoch hợp lệ tại mốc đó; số seed n được ghi trong bảng và điểm n<5 được chú thích trên hình. Thanh sai số là CI 95% t với df=n−1. Tiêu chí 0.5 điểm phần trăm cung cấp một quy tắc rõ ràng để cân bằng độ chính xác với thời gian; nó không thay thế đánh giá chi phí năng lượng hoặc giá trị triển khai.

{selection_text} Khi checkpoint sớm tiếp tục được chọn ở mốc dài, đường best-so-far phẳng chỉ cho thấy fine-tuning chưa vượt checkpoint đó. Cần đối chiếu đường epoch cuối để nhận biết giảm hiệu năng rồi phục hồi; không diễn giải riêng đường best-so-far thành hội tụ ổn định.

### So sánh YOLOv13 với baseline

YOLOv13n và YOLOv13s đạt sweet spot lần lượt ở {sweet['yolov13n']['sweet_budget']} và {sweet['yolov13s']['sweet_budget']} phút. YOLO11n và YOLOv12n đạt ở {sweet['yolo11n']['sweet_budget']} và {sweet['yolov12n']['sweet_budget']} phút. Theo test mAP@0.5:0.95 tại checkpoint sweet spot, mô hình cao nhất là {MODEL_NAMES[best_test_model]} ({best_test['mean'] * 100:.2f}% trung bình); phải so sánh CI và độ lệch seed trước khi kết luận khác biệt có ý nghĩa. Thí nghiệm này không khẳng định YOLOv13 luôn vượt baseline; lựa chọn mô hình cần cân nhắc đồng thời latency, kích thước và hiệu năng.

{stability_text} Chênh lệch ghép theo seed dưới đây mang tính mô tả, chưa hiệu chỉnh đa so sánh hoặc độ bất định do chọn mốc; ngân sách của từng model có thể khác nhau:

{chr(10).join(comparison_lines)}

### Hội tụ, overfitting và giới hạn

Loss và mAP epoch cuối lấy từ epoch hoàn tất cuối cùng trong ngân sách; mAP best-so-far lấy từ checkpoint có validation tốt nhất trong toàn bộ thời gian đã chạy đến mốc đó. Hai chuỗi mô tả hai góc nhìn khác nhau và được đặt cạnh nhau ở bảng cùng biểu đồ. Bảng dưới đếm seed có mAP epoch cuối còn tăng hơn 0.5 điểm trong 60→120 phút và seed có dấu hiệu overfit ứng viên (train loss giảm ≥5%, val loss tăng ≥5%, còn mAP epoch cuối từ 5→120 phút tăng không quá 0.5 điểm). Chỉ ghép các seed đủ mốc 5, 60 và 120 phút. Đây là chỉ báo mô tả, không phải kiểm định thống kê; không kết luận underfitting chỉ từ loss còn cao.

{fit_table}

Năm seed đo biến thiên ngẫu nhiên của huấn luyện trên một split cố định, không đo độ bất định do cách chia dữ liệu.

**Internal validity:** {pretrain_source_test:,} ảnh test bắt nguồn từ COCO train2017, có nguy cơ đã xuất hiện trong pretraining; vì vậy holdout chỉ độc lập với fine-tuning. Dữ liệu được chia theo ảnh trong từng stratum, chưa chứng minh tách biệt người hoặc chuỗi video AdvT-shirt. Cùng hyperparameters chưa chắc tối ưu cho mỗi kiến trúc; pretrained khác nhau, tốc độ epoch khác nhau và chọn checkpoint trên validation cùng ảnh hưởng so sánh. Không suy ra lợi ích kiến trúc thuần túy từ benchmark này.

**External validity:** COCO chỉ là mẫu 10.000 ảnh, không phải toàn bộ benchmark gốc; bản AdvT-shirt mirror thiếu 8 ảnh so với trang dự án. Target chỉ gồm person, không đo phát hiện patch hay khả năng chống tấn công. Một split cố định và năm seed chưa đủ bảo đảm triển khai rộng; cần test ngoài miền và tách theo người/chuỗi, đồng thời loại trừ ảnh đã dùng để pretrain khi đánh giá khả năng tổng quát hóa.

**Hardware validity:** RTX 5060 không đại diện cho thiết bị edge đích. Latency đo một lượt cho mỗi seed theo thứ tự cố định, chịu ảnh hưởng của nhiệt độ, xung GPU, driver, CUDA và pipeline; CI qua năm checkpoint không thay thế phép đo lặp có thứ tự ngẫu nhiên. FPS nghịch đảo median không phải throughput camera hay throughput xử lý đồng thời. Không đo năng lượng; cần benchmark lại đúng thiết bị, định dạng model và pipeline ứng dụng cuối.

## 5. Khuyến nghị triển khai và kết luận

{accuracy_choice_text}

{realtime_text}

Khuyến nghị chọn ngân sách huấn luyện theo sweet spot đã đo cho đúng model, dữ liệu và mục tiêu triển khai; dùng seed bổ sung và split ngoài miền trước khi khóa lựa chọn production. Nếu ngân sách cho phép, giữ đánh giá ở 120 phút như một kiểm tra trần hiệu năng; dừng sớm ở sweet spot khi lợi ích thêm thấp hơn ngưỡng 0.5 điểm phần trăm. Trước khi triển khai edge, đo lại latency trên thiết bị đích, batch/decode pipeline thực và định dạng xuất cuối.

{boundary_text} Với model liên tục chọn checkpoint rất sớm, cần đo thêm trọng số pretrained ở mốc 0 phút và điều chỉnh learning rate/warm-up để xác định fine-tuning có đem lại lợi ích. Sweet spot pilot và nghiên cứu mở rộng không thể dùng thay nhau khi kích thước dữ liệu, tốc độ epoch và cách chọn checkpoint thay đổi.

Nghiên cứu cung cấp pipeline tái lập cho dữ liệu person-detection đã hợp nhất, bốn mô hình, năm seed, sáu ngân sách thời gian, test holdout, profiling batch 1, và quy tắc sweet spot định trước. Kết luận về mô hình tốt nhất chỉ rút ra từ kết quả thực đo trong bảng, với giới hạn được nêu rõ.

## Tài liệu tham khảo và dữ liệu

1. COCO 2017 Detection: <https://cocodataset.org/dataset/detection-2017.htm>.
2. AdvT-shirt-1K source project: <https://github.com/Wwangb/AdvT-shirt-1K>; dataset card mirror: <https://huggingface.co/datasets/xingjunm/AdvT-shirt>.
3. M. Lei et al., “YOLOv13: Real-Time Object Detection with Hypergraph-Enhanced Adaptive Visual Perception,” arXiv:2506.17733, 2025. <https://arxiv.org/abs/2506.17733>; official repository: <https://github.com/iMoonLab/yolov13>.
4. YOLOv12 official repository: <https://github.com/sunsmarterjie/yolov12>.
5. YOLO11 official Ultralytics documentation: <https://docs.ultralytics.com/models/yolo11/>.

## Mã LaTeX IEEE cho các bảng

### Bảng 1

```latex
{latex1}
```

### Bảng 2

```latex
{latex2}
```
"""
    return markdown, sweet, curves, latency, diagnostics


def material_passport(body: str, records: dict, sweet: dict, latency: dict, md_size: int, docx_size: int) -> str:
    state = read_json(project_path("results/fullscale_experiments/orchestrator_state.json", must_exist=True))
    starts = [datetime.fromisoformat(records[m][s]["started_at_utc"]) for m in MODELS for s in SEEDS]
    ends = [datetime.fromisoformat(latency[m][s]["measured_at_utc"]) for m in MODELS for s in SEEDS]
    duration = int((max(ends) - min(starts)).total_seconds())
    audit = read_json(project_path("results/fullscale_setup/fullscale_dataset_audit.json", must_exist=True))
    audit_issues = sum(split["invalid_label_rows"] + len(split["missing_label_stems"]) + len(split["orphan_label_stems"]) for split in audit["splits"].values())
    manifest = read_json(project_path("dataset/fullscale_person/manifest.json", must_exist=True))
    sanitization = manifest.get("advtshirt_box_sanitization", {})
    clipped = sanitization.get("source_out_of_bounds_rows", 0)
    reencoded = sanitization.get("rows_reencoded_or_edge_adjusted", 0)
    max_overrun = sanitization.get("maximum_source_overrun_normalized", 0.0)
    out_files = [
        ("results/BaoCao_TrienKhaiRong_YOLOv13.md", md_size),
        ("results/BaoCao_TrienKhaiRong_YOLOv13.docx", docx_size),
        ("results/fullscale_mAP_growth.png", (OUT / "fullscale_mAP_growth.png").stat().st_size),
        ("results/fullscale_experiments/latency_records.json", (EXP / "latency_records.json").stat().st_size),
        ("results/fullscale_experiments/epoch_history/ (20 CSVs)", sum(p.stat().st_size for p in (EXP / "epoch_history").glob("*/*.csv"))),
    ]
    if state.get("status") != "complete" or any(state.get("jobs", [])[i].get("status") != "complete" for i in range(len(state.get("jobs", [])))):
        raise RuntimeError("Cannot publish a completed experiment passport while orchestrator jobs remain incomplete")
    spots = ", ".join(f"{MODEL_NAMES[m]} {sweet[m]['sweet_budget']}m" for m in MODELS)
    missing_milestones = [f"{MODEL_NAMES[m]} seed {s} at {row['budget_minutes']}m" for m in MODELS for s in SEEDS for row in records[m][s]["milestones"] if row.get("status") != "complete"]
    anomalies = f"{clipped} source AdvT-shirt boxes exceeded image bounds (max overrun {max_overrun:.8g} normalized); {reencoded} rows were re-encoded with edge headroom; final audit issues={audit_issues}."
    if missing_milestones:
        anomalies += " No completed epoch within budget for: " + "; ".join(missing_milestones) + ". These milestones are excluded from sweet-spot selection."
    return f"""## Material Passport

- Origin Skill: experiment-agent
- Origin Mode: run
- Origin Date: {REPORT_CREATED_AT.isoformat()}
- Verification Status: UNVERIFIED (multi-seed run completed; independent full rerun not performed)
- Version Label: exp_result_v1

## Experiment Result

- **ID**: fullscale-20261011
- **Type**: training
- **Status**: completed
- **Command**: `.venv\\Scripts\\python.exe run_fullscale_multiseed.py`; `.venv\\Scripts\\python.exe benchmark_fullscale_latency.py --all`; `.venv\\Scripts\\python.exe generate_fullscale_report.py`
- **Working Directory**: `{PROJECT_ROOT}`
- **Duration**: {duration} seconds from first training start through final latency benchmark
- **Exit Code**: {state.get('exit_code', 0)}

### Output Files

| File | Size |
|------|------:|
{chr(10).join(f'| `{name}` | {size:,} bytes |' for name, size in out_files)}

### Output Summary

Four models, five seeds each, six time budgets; validation-selected sweet spots: {spots}. Per-epoch training histories are archived with SHA-256 checksums.

### Anomalies Detected

{anomalies}

"""


def _shade(cell, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _set_repeat_table_header(row) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(repeat)


def add_docx(markdown: str, records: dict, sweet: dict, curves: dict, latency: dict, diagnostics: dict) -> Path:
    from docx import Document
    from docx.enum.section import WD_ORIENT, WD_SECTION_START
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml.ns import qn

    doc = Document()
    doc.core_properties.author = "Mai Quốc Bảo"
    doc.core_properties.last_modified_by = "Mai Quốc Bảo"
    doc.core_properties.title = "Báo cáo thực nghiệm quy mô lớn về hội tụ thời gian và hiệu quả triển khai YOLOv13"
    doc.core_properties.language = "vi-VN"
    doc.core_properties.created = REPORT_CREATED_AT
    doc.core_properties.modified = REPORT_CREATED_AT
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(7)
    for name, size in (("Title", 18), ("Heading 1", 15), ("Heading 2", 13), ("Heading 3", 12)):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.line_spacing = 1.15
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            style._element.rPr.rFonts.attrib.pop(qn(f"w:{attr}"), None)
        style._element.rPr.rFonts.set(qn("w:cs"), "Times New Roman")
        if style._element.pPr is not None:
            for border in list(style._element.pPr.findall(qn("w:pBdr"))):
                style._element.pPr.remove(border)

    def clean_text(text):
        text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
        text = re.sub(r"<(https?://[^>]+)>", r"\1", text)
        return text.replace("**", "").replace("`", "")

    def bullet(text):
        p = doc.add_paragraph(clean_text(text), style="List Bullet")
        p.paragraph_format.line_spacing = 1.15
        return p

    def prose_block(text):
        lines = text.splitlines()
        if lines and all(line.startswith("- ") for line in lines):
            for line in lines:
                bullet(line[2:])
        else:
            para(text)

    def para(text="", bold_lead=None):
        text = clean_text(text)
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.15
        if bold_lead and text.startswith(bold_lead):
            p.add_run(bold_lead).bold = True
            p.add_run(text[len(bold_lead):])
        else:
            p.add_run(text)
        return p

    passport_text = markdown.split("## Material Passport\n\n", 1)[1].split("\n\n## Experiment Result", 1)[0]
    result_text = markdown.split("## Experiment Result\n\n", 1)[1].split("\n\n# Báo cáo", 1)[0]

    title = doc.add_paragraph(style="Title")
    title.alignment = 1
    title.add_run("Báo cáo thực nghiệm quy mô lớn về hội tụ thời gian và hiệu quả triển khai YOLOv13")
    subtitle = doc.add_paragraph()
    subtitle.alignment = 1
    subtitle.add_run("Bốn mô hình với năm seed và sáu ngân sách huấn luyện").italic = True
    para(f"Tác giả: Mai Quốc Bảo · Ngày lập: {datetime.now().astimezone().strftime('%d-%m-%Y')}")

    def orientation(landscape: bool):
        new = doc.add_section(WD_SECTION_START.NEW_PAGE)
        new.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
        new.page_width, new.page_height = (Inches(11), Inches(8.5)) if landscape else (Inches(8.5), Inches(11))
        new.top_margin = new.bottom_margin = Inches(0.6)
        new.left_margin = new.right_margin = Inches(0.7)

    def table_layout(table, widths):
        from docx.oxml import OxmlElement
        table.autofit = False
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            border = OxmlElement(f"w:{side}")
            for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                border.set(qn(f"w:{key}"), value)
            borders.append(border)
        table._tbl.tblPr.append(borders)
        for column, width in zip(table.columns, widths):
            column.width = Inches(width)
        for index, row in enumerate(table.rows):
            if index == 0:
                _set_repeat_table_header(row)
            no_split = OxmlElement("w:cantSplit")
            row._tr.get_or_add_trPr().append(no_split)
            for cell, width in zip(row.cells, widths):
                cell.width = Inches(width)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if index:
                    _shade(cell, "F1F5FA" if index % 2 == 0 else "FFFFFF")
                margins = OxmlElement("w:tcMar")
                for side in ("top", "left", "bottom", "right"):
                    margin = OxmlElement(f"w:{side}")
                    margin.set(qn("w:w"), "80")
                    margin.set(qn("w:type"), "dxa")
                    margins.append(margin)
                cell._tc.get_or_add_tcPr().append(margins)
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing = 1.15
                    p.paragraph_format.keep_with_next = index == 0
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.space_after = Pt(3)
                    for run in p.runs:
                        run.font.name = "Times New Roman"
                        run.font.size = Pt(12)

    doc.add_heading("Tóm tắt", 1)
    abstract = markdown.split("## Tóm tắt\n\n", 1)[1].split("\n\n## 1.", 1)[0]
    for block in abstract.split("\n\n"):
        if block.startswith("- **"):
            for line in block.splitlines():
                bullet(line[2:] if line.startswith("- ") else line)
        else:
            para(block.replace("**", ""))

    doc.add_heading("1 Thiết lập thực nghiệm", 1)
    for h, content in (("Dữ liệu và nhãn", markdown.split("### Dữ liệu và nhãn\n\n", 1)[1].split("\n\n### Mô hình", 1)[0]),
                       ("Mô hình và điều kiện huấn luyện", markdown.split("### Mô hình và điều kiện huấn luyện\n\n", 1)[1].split("\n\n### Chỉ số", 1)[0]),
                       ("Chỉ số và thống kê", markdown.split("### Chỉ số và thống kê\n\n", 1)[1].split("\n\n## 2.", 1)[0])):
        doc.add_heading(h.replace(",", ""), 2)
        for line in content.splitlines():
            line = line.strip()
            if line.startswith("-"):
                bullet(line[2:] if line.startswith("- ") else line)
            elif line:
                para(line.replace("**", "").replace("`", ""))

    orientation(True)
    doc.add_heading("2 Bảng so sánh đa seed tại sweet spot", 1)
    para("Checkpoint được chọn bằng validation; các chỉ số phát hiện được đánh giá trên test holdout đối với fine-tuning. Bảng trình bày μ ± σ và khoảng tin cậy 95% của trung bình.")
    headers = ["Mô hình", "Params\n(M)", "FLOPs\n(G)", "Precision\n(%)", "Recall\n(%)", "mAP@0.5\n(%)", "mAP@0.5:0.95\n(%)", "Latency\n(ms/frame)", "FPS"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_repeat_table_header(table.rows[0])
    for cell, label in zip(table.rows[0].cells, headers):
        cell.text = label
        _shade(cell, "17365D")
        for p in cell.paragraphs:
            p.alignment = 1
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(9)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
    for model in MODELS:
        budget = sweet[model]["sweet_budget"]
        rows = [milestone_for(records[model][seed], budget) for seed in SEEDS]
        metrics = {key: stats([r["test"][key] for r in rows if r and r.get("test") and r["test"].get(key) is not None]) for key in ("precision", "recall", "map50", "map50_95")}
        pstat = stats([r["parameters_m"] for r in rows if r and r.get("parameters_m") is not None])
        fstat = stats([r["flops_g"] for r in rows if r and r.get("flops_g") is not None])
        lstat = stats([latency[model][s]["latency_median_ms_per_frame"] for s in SEEDS])
        fpsstat = stats([latency[model][s]["fps_from_median_latency"] for s in SEEDS])
        vals = [MODEL_NAMES[model], fmt(pstat["mean"], 3), fmt(fstat["mean"], 2)]
        vals.extend(fmt_stats(metrics[k], 100, 1).replace("; 95% CI ", "\nCI ") for k in ("precision", "recall", "map50", "map50_95"))
        vals.extend((fmt_stats(lstat, 1, 1, "ms").replace("; 95% CI ", "\nCI "), fmt_stats(fpsstat, 1, 1).replace("; 95% CI ", "\nCI ")))
        cells = table.add_row().cells
        for idx, (cell, value) in enumerate(zip(cells, vals)):
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.alignment = 1 if idx else 0
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(8)
    para("CI là khoảng tin cậy t Student 95% của trung bình; μ ± σ là trung bình và độ lệch chuẩn mẫu qua năm seed.")
    table_layout(table, [1.0, 0.75, 0.65, 1.15, 1.15, 1.15, 1.4, 1.2, 1.15])

    doc.add_heading("3 Hội tụ theo thời gian và điểm tối ưu", 1)
    para("Bảng tách mAP95 của checkpoint validation tốt nhất trong ngân sách và mAP95 của epoch cuối hoàn tất. Δ và lợi ích biên được tính theo best-so-far chỉ khi cả hai mốc liền trước và hiện tại đủ cùng năm seed; loss là tổng box + class + DFL tại epoch cuối dưới ngân sách.")
    headers2 = ["Mô hình", "Phút", "n", "Best-checkpoint val mAP95 (%)\nμ±σ; 95% CI", "Last-epoch val mAP95 (%)\nμ±σ; 95% CI", "Δ best (pp)", "Gain best\n(pp/min)", "Train loss", "Val loss"]
    table2 = doc.add_table(rows=1, cols=len(headers2))
    table2.style = "Table Grid"
    table2.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_repeat_table_header(table2.rows[0])
    for cell, label in zip(table2.rows[0].cells, headers2):
        cell.text = label
        _shade(cell, "17365D")
        for p in cell.paragraphs:
            p.alignment = 1
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(9)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
    for model in MODELS:
        previous_budget = None
        previous_mean = None
        previous_n = None
        for budget in BUDGETS:
            cell_data = curves[model][budget]
            summary = cell_data["map"]
            chosen = sweet[model]["sweet_budget"] == budget
            if not summary["n"]:
                vals = [MODEL_NAMES[model], str(budget), "0", "—", "—", "—", "—", "—", "—"]
            else:
                delta = (summary["mean"] - previous_mean) * 100 if previous_mean is not None and summary["n"] == previous_n == len(SEEDS) else None
                dt = budget - previous_budget if previous_budget is not None else None
                vals = [MODEL_NAMES[model], str(budget) + (" ★" if chosen else ""), str(summary["n"]), fmt_stats(summary, 100, 1).replace("; 95% CI ", "\nCI "), fmt_stats(cell_data["endpoint_map"], 100, 1).replace("; 95% CI ", "\nCI "), fmt(delta), fmt(delta / dt if delta is not None and dt else None), fmt(cell_data["train_loss"]["mean"]), fmt(cell_data["val_loss"]["mean"])]
                previous_budget, previous_mean = budget, summary["mean"]
                previous_n = summary["n"]
            cells = table2.add_row().cells
            for idx, (cell, value) in enumerate(zip(cells, vals)):
                cell.text = value
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for p in cell.paragraphs:
                    p.alignment = 1 if idx else 0
                    p.paragraph_format.line_spacing = 1.0
                    p.paragraph_format.space_after = Pt(0)
                    for run in p.runs:
                        run.font.name = "Times New Roman"
                        run.font.size = Pt(8)
                        if chosen and idx in (0, 1):
                            run.font.bold = True

    table_layout(table2, [1.0, 0.5, 0.3, 2.2, 2.2, 0.8, 1.0, 0.8, 0.8])
    doc.add_heading("Tín hiệu loss và tiếp tục hội tụ", 2)
    para("Ứng viên overfit yêu cầu train loss giảm ít nhất 5%, val loss tăng ít nhất 5% và mAP validation tại epoch cuối từ 5 đến 120 phút tăng không quá 0.5 điểm. Mức tăng mAP epoch cuối 60→120 phút trên 0.5 điểm cho thấy vẫn còn cải thiện đo được. Chỉ ghép seed đủ cả ba mốc.")
    diag_table = doc.add_table(rows=1, cols=7)
    diag_table.style = "Table Grid"
    diag_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, label in zip(diag_table.rows[0].cells, ["Mô hình", "Seed ghép cặp", "Δ train loss", "Δ val loss", "Δ mAP 60→120 (pp)", "Seed còn tăng", "Ứng viên overfit"]):
        cell.text = label
        _shade(cell, "D9EAF7")
    for model in MODELS:
        d = diagnostics[model]
        values = [MODEL_NAMES[model], str(d["paired_seeds"]), fmt(d["mean_train_loss_change_5_to_120"]["mean"]), fmt(d["mean_val_loss_change_5_to_120"]["mean"]), fmt(d["mean_val_map_gain_60_to_120_pp"]["mean"]), str(d["continued_learning_seeds"]), str(d["overfit_candidate_seeds"])]
        for cell, value in zip(diag_table.add_row().cells, values):
            cell.text = value
    table_layout(diag_table, [1.2, 1.0, 1.3, 1.3, 2.0, 1.3, 1.5])

    doc.add_heading("Đường cong tăng trưởng mAP", 2)
    doc.add_picture(str(OUT / "fullscale_mAP_growth.png"), width=Inches(8.6))
    cap = doc.add_paragraph("Hình 1. Trái: mAP95 của checkpoint validation tốt nhất trong ngân sách. Phải: mAP95 của epoch cuối đã hoàn tất. Thanh sai số là CI 95% t Student đa seed.")
    cap.alignment = 1
    for r in cap.runs:
        r.italic = True
        r.font.size = Pt(10)

    orientation(False)
    doc.add_heading("Ma trận sweet spot", 2)
    matrix = doc.add_table(rows=1, cols=6)
    matrix.style = "Table Grid"
    matrix.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, label in zip(matrix.rows[0].cells, ["Mô hình", "Sweet spot", "Peak quan sát", "Val mAP95 (%)\nMean [95% CI]", "Khoảng cách đến peak (pp)", "Số seed"]):
        cell.text = label
        _shade(cell, "D9EAF7")
    for model in MODELS:
        item = sweet[model]
        ci = item["sweet_stats"]["ci95"]
        vals = [MODEL_NAMES[model], f"{item['sweet_budget']} phút", f"{item['peak_budget']} phút", f"{item['sweet_mean']*100:.2f} [{ci[0]*100:.2f}, {ci[1]*100:.2f}]", f"{item['extra_gain_pp']:.2f}", "5"]
        for cell, value in zip(matrix.add_row().cells, vals):
            cell.text = value
    table_layout(matrix, [1.0, 0.8, 0.8, 2.1, 1.4, 1.0])
    doc.add_heading("4 Thảo luận", 1)
    for h, content in (("Tăng trưởng và diminishing returns", markdown.split("### Tăng trưởng và diminishing returns\n\n", 1)[1].split("\n\n### So sánh", 1)[0]),
                       ("So sánh YOLOv13 với baseline", markdown.split("### So sánh YOLOv13 với baseline\n\n", 1)[1].split("\n\n### Hội tụ", 1)[0]),
                       ("Hội tụ, overfitting và giới hạn", markdown.split("### Hội tụ, overfitting và giới hạn\n\n", 1)[1].split("\n\n## 5.", 1)[0])):
        doc.add_heading(h.replace(",", ""), 2)
        if h == "Hội tụ, overfitting và giới hạn" and "\n\n| Model | Paired seeds |" in content:
            before, after = content.split("\n\n| Model | Paired seeds |", 1)
            content = before + ("\n\n" + after.split("\n\n", 1)[1] if "\n\n" in after else "")
        for block in content.split("\n\n"):
            prose_block(block)
    doc.add_heading("5 Khuyến nghị triển khai và kết luận", 1)
    conclusion = markdown.split("## 5. Khuyến nghị triển khai và kết luận\n\n", 1)[1].split("\n\n## Tài liệu tham khảo", 1)[0]
    for block in conclusion.split("\n\n"):
        prose_block(block)
    doc.add_heading("Tài liệu tham khảo và dữ liệu", 1)
    refs = markdown.split("## Tài liệu tham khảo và dữ liệu\n\n", 1)[1].split("\n\n## Mã LaTeX", 1)[0]
    for line in refs.splitlines():
        if line.strip():
            doc.add_paragraph(clean_text(re.sub(r"^\d+\.\s*", "", line.strip())), style="List Number")

    doc.add_page_break()
    doc.add_heading("Phụ lục nhật ký và bằng chứng thực nghiệm", 1)
    para("Mã nguồn bảng LaTeX IEEE được cung cấp trong báo cáo Markdown đi kèm. Checkpoint và dữ liệu thô được lưu cục bộ; bản ghi thống kê và CSV theo epoch được lưu trong results/fullscale_experiments.")
    for line in (passport_text + "\n" + result_text).splitlines():
        if line.strip() and not line.startswith("|") and not line.startswith("###"):
            para(line.strip().lstrip("- "))

    footer = section.footer.paragraphs[0]
    footer.alignment = 1
    footer.add_run("Báo cáo thực nghiệm YOLO · ")
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    for run in footer.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(9)

    target = OUT / "BaoCao_TrienKhaiRong_YOLOv13.docx"
    doc.save(target)
    return target


def main() -> int:
    require_local_venv()
    from validate_fullscale_results import audit_results
    audit_results(require_complete=True)
    records = load_records()
    archive_epoch_histories(records)
    markdown, sweet, curves, latency, diagnostics = make_report(records)
    md_path = OUT / "BaoCao_TrienKhaiRong_YOLOv13.md"
    docx_path = OUT / "BaoCao_TrienKhaiRong_YOLOv13.docx"
    for _ in range(4):
        passport = material_passport(markdown, records, sweet, latency,
                                     md_path.stat().st_size if md_path.exists() else 0,
                                     docx_path.stat().st_size if docx_path.exists() else 0)
        complete_markdown = passport + markdown
        md_path.write_text(complete_markdown, encoding="utf-8")
        add_docx(complete_markdown, records, sweet, curves, latency, diagnostics)
    complete_markdown = material_passport(markdown, records, sweet, latency, md_path.stat().st_size, docx_path.stat().st_size) + markdown
    md_path.write_text(complete_markdown, encoding="utf-8")
    add_docx(complete_markdown, records, sweet, curves, latency, diagnostics)
    print(json.dumps({"markdown": str(md_path.relative_to(PROJECT_ROOT)), "docx": str(docx_path.relative_to(PROJECT_ROOT)), "sweet_spots_minutes": {MODEL_NAMES[m]: sweet[m]["sweet_budget"] for m in MODELS}}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
