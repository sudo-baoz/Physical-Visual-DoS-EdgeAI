"""Author the research update from completed and audited experimental evidence."""
from __future__ import annotations
import argparse
import csv
import json
import math
from datetime import datetime, timezone, timedelta
from pathlib import Path
from pipeline_utils import configure_project_environment,project_path,require_local_venv,sha256_file,write_json

TITLE="Tối ưu giao thức chuyển giao và đối sánh khoa học YOLOv13 cho phát hiện người"
STEM="BaoCao_NCKH_YOLOv13_Updated"
OUT=project_path("results")
EDITOR=project_path("results/updated_research")
METRICS=("precision","recall","map50","map50_95")
LATENCY_PROTOCOL="actual-batch1-txt-loader-v2"


def read(path):
    return json.loads(project_path(path,must_exist=True).read_text(encoding="utf-8"))


def fmt(stat,scale=1.,decimals=2,ci=True):
    text=f"{float(stat['mean'])*scale:.{decimals}f}"
    if stat.get("sd") is not None:text+=f" ± {float(stat['sd'])*scale:.{decimals}f}"
    lo,hi=stat.get("ci95_low"),stat.get("ci95_high")
    if lo is None and stat.get("ci95") is not None:lo,hi=stat["ci95"]
    if ci and lo is not None:text+=f" [{lo*scale:.{decimals}f}; {hi*scale:.{decimals}f}]"
    return text


def require_inputs():
    data=read("results/optimization/final_summary.json")
    lit=read("results/literature_updated/benchmark_matrix.json")
    figures=read("results/figures/figures_manifest.json")
    if data.get("status")!="complete" or figures.get("status")!="complete":raise RuntimeError("Experiments and figures must be complete")
    if read("results/literature_updated/literature_audit.json").get("status")!="pass":raise RuntimeError("Literature audit must pass")
    audit=read("results/updated_research/metric_integrity_audit.json")
    if audit.get("status") not in ("pass","passed"):raise RuntimeError("Independent metric integrity audit must pass")
    if audit.get("final_summary_sha256")!=sha256_file("results/optimization/final_summary.json"):
        raise RuntimeError("Metric audit is stale for final_summary")
    for path,expected in figures.get("input_sha256",{}).items():
        if sha256_file(path)!=expected:raise RuntimeError(f"Figure input changed: {path}")
    for artifact in figures.get("artifacts",[]):
        if sha256_file(artifact["path"])!=artifact["sha256"]:raise RuntimeError(f"Figure changed: {artifact['path']}")
    if data["winner_lock"].get("test_metrics_seen_before_lock") is not False:raise RuntimeError("Test selection gate")
    if data.get("pretrained_locked_test_diagnostic",{}).get("status")!="complete":raise RuntimeError("Native pretrained locked test diagnostic required")
    def valid_latency(lat):
        if lat.get("protocol_id")!=LATENCY_PROTOCOL or lat.get("actual_batch_size")!=1:
            raise RuntimeError("Legacy/unverified latency must not enter final report")
        if (lat.get("warmup_frames"),lat.get("measured_frames"))!=(32,500):
            raise RuntimeError("Final latency requires full 32+500 measurement, not a probe")
        if set(lat.get("network_forward_batch_size_counts",{}))!={"1"}:
            raise RuntimeError("Network-forward batch size must be verified as one")
    valid_latency(data["pretrained_locked_test_diagnostic"]["latency"])
    if data.get("zero_locked_test_diagnostic"):valid_latency(data["zero_locked_test_diagnostic"]["latency"])
    for point in figures.get("local_operating_points",[]):
        if point.get("kind")=="original_multiseed" and (point.get("latency_ms",{}).get("mean") is not None or point.get("fps",{}).get("mean") is not None):
            raise RuntimeError("Historical timing cohort excluded because of unexplained runtime-context drift")
    seeds=set(data["protocol"]["confirmation_seeds"])
    for g in data["groups"]:
        if g["n"]!=3 or {r["seed"] for r in g["seed_records"]}!=seeds:raise RuntimeError("Incomplete confirmation")
        for k in (*METRICS,"parameters_m","flops_g","latency_ms","fps"):
            s=g["metrics"][k]
            if s["n"]!=3 or not all(math.isfinite(float(s[x])) for x in ("mean","sd","ci95_low","ci95_high")):raise RuntimeError(f"Invalid {k}")
        for r in g["seed_records"]:
            project_path(r["checkpoint"],must_exist=True);valid_latency(r["latency"])
    return data,lit,figures


def blocks_for(data,lit,figures):
    blocks=[]
    def h(text,level=1):blocks.append(dict(kind="heading",text=text,level=level))
    def p(text,landscape=None):blocks.append(dict(kind="paragraph",text=text,landscape=landscape))
    def t(caption,headers,rows,widths,landscape=False,latex=False):
        blocks.append(dict(kind="table",caption=caption,headers=headers,rows=rows,widths=widths,landscape=landscape,latex=latex))
    control=next(g for g in data["groups"] if g["configuration"]=="original_sgd")
    proposed=next(g for g in data["groups"] if g["configuration"]!="original_sgd")
    delta=data["paired_test_map95_delta"];delta_text=fmt(delta,100)
    verdict="CI ghép cặp nằm trên 0, hỗ trợ cải thiện trong ba seed đã đo." if delta["ci95_low"]>0 else "CI ghép cặp không nằm hoàn toàn trên 0; chưa xác nhận cải thiện ổn định trong ba seed đã đo."
    pre=data["pretrained_validation"]["validation"];winner=data["winner_configuration"]
    native=data["pretrained_locked_test_diagnostic"]
    native_test=native["test"];native_latency=native["latency"]
    deployment=data["deployment_choice"];zero_test=data.get("zero_locked_test_diagnostic")
    day=datetime.now(timezone(timedelta(hours=7))).strftime("%d/%m/%Y")
    p(f"Tác giả: Mai Quốc Bảo. Báo cáo gửi Hội đồng và giảng viên hướng dẫn. Cập nhật: {day}.")
    h("Tóm tắt")
    p(f"Nghiên cứu kiểm tra cải thiện giao thức chuyển giao và fine tuning YOLOv13s cho person trên 11.123 ảnh. Sau sàng lọc {len(data['screening_records'])} cấu hình bằng validation với seed 42, cấu hình {winner} được khóa và đối chiếu SGD ban đầu ở ba seed 100, 2026 và 777. Mỗi quỹ đạo gồm sáu epoch, ảnh 416, batch 8; ba đối chứng dùng prefix lịch sử đã audit và ba nhánh tối ưu được huấn luyện mới. Test mAP@0.5:0.95 của đối chứng là {fmt(control['metrics']['map50_95'],100)}%, của nhánh tối ưu là {fmt(proposed['metrics']['map50_95'],100)}%; chênh lệch ghép cặp là {delta_text} điểm phần trăm. {verdict}")
    if zero_test and deployment["choice"]=="zero_restored_person":
        p(f"Theo lựa chọn triển khai đã khóa chỉ bằng validation, checkpoint một lớp khôi phục từ pretrained và chưa fine tune được ưu tiên. Điểm này đạt test AP95 {zero_test['test']['map50_95']*100:.2f}%, latency {zero_test['latency']['latency_median_ms_per_frame']:.2f} ms/frame, n=1 không CI đa seed. Lựa chọn tiết kiệm huấn luyện bổ sung chỉ được hỗ trợ trong split và giao thức đã đo; vẫn cần kiểm tra trên dữ liệu mới.")
    else:
        p(f"Lựa chọn triển khai được khóa từ validation seed 42 là nhánh đã train {winner}. Quyết định này tách khỏi đánh giá test; không biến số test thành tiêu chí chọn cấu hình.")
        if zero_test:
            p(f"Checkpoint khôi phục một lớp chưa fine tune được đo thêm như điểm chẩn đoán n=1: test AP95 {zero_test['test']['map50_95']*100:.2f}%, latency {zero_test['latency']['latency_median_ms_per_frame']:.2f} ms/frame, không CI đa seed. Điểm này không phải lựa chọn triển khai đã khóa và kết quả test không được dùng đổi quyết định.")
    p(f"Latency trung bình của median từng seed nhánh tối ưu là {fmt(proposed['metrics']['latency_ms'])} ms/frame, PyTorch FP32 batch1 trên RTX5060. Diagnostic checkpoint COCO giữ head80 và lọc person cho validation AP95 {pre['map50_95']*100:.2f}%. Audit tensor phát hiện cấu trúc YOLOv13s tái dựng trước khác checkpoint pretrained, cần khai báo khi diễn giải kết quả cũ. Đóng góp là kiểm tra chuyển giao, thích ứng giao thức và đánh giá có đối chứng; chưa là kiến trúc mới, tối ưu toàn cục, SOTA hoặc phòng thủ mọi patch. CI phản ánh biến thiên seed trên một split cố định; test đã được xem trong nghiên cứu trước.")
    p(f"Đối chiếu bổ sung sau khóa với checkpoint native không fine tune, giữ head 80 lớp và lọc person, cho test AP95 {native_test['map50_95']*100:.2f}% tại một điểm cố định, không CI đa seed. Vì vậy cải thiện so với SGD cũ không tự chứng minh fine tuning tốt hơn checkpoint pretrained. Các nguồn COCO trong test có thể từng nằm trong pretraining.")
    p("Từ khóa: YOLOv13; chuyển giao; HyperACE; phát hiện người; đối sánh giao thức; đa seed.")
    h("1 Đặt vấn đề và câu hỏi nghiên cứu")
    p("Nghiên cứu trước đạt test AP95 47,97% ở cấu hình được gọi YOLOv13s và sweet spot 10 phút theo checkpoint validation tốt nhất. Giá trị này chưa giải thích chất lượng chuyển giao hay khả năng chống patch. Cập nhật kiểm tra tensor, scheduler, loss, augmentation và sampling, rồi đo hiệu quả và chi phí bằng đối chứng mới.")
    p("Câu hỏi chính là cấu hình chọn chỉ bằng validation có cải thiện test AP so với SGD trên cùng dữ liệu và sáu epoch hay không. Câu hỏi bổ sung là cải thiện thuộc nguồn COCO hay AdvT-shirt, và chi phí phù hợp triển khai không. Nếu nhánh sửa cấu trúc pretrained thắng, phép so sánh đo tác dụng gộp sửa chuyển giao và recipe; chưa phân biệt nhân quả từng yếu tố.")
    h("2 Nghiên cứu liên quan và khác biệt giao thức")
    p("Rà soát có chủ đích các nguồn gốc liên quan, chưa là tổng quan hệ thống toàn bộ SOTA. YOLO11 [3] được dẫn như phần mềm chính thức; YOLOv12 [2], YOLOv13 [1], YOLOv10 [4], RT-DETR [5] và RT-DETRv2 [6] là các công trình kiến trúc. Bảng I giữ riêng số công bố. NR nghĩa là chưa xác định trong nguồn được kiểm tra.")
    ref_index={r["id"]:i+1 for i,r in enumerate(lit["references"])}
    rows=[]
    for r in lit["rows"]:
        def number(k):return "NR" if r.get(k) is None else f"{r[k]:g}"
        speed=f"{r['published_latency_ms']:g} ms" if r.get("published_latency_ms") is not None else f"{r['published_fps']:g} FPS" if r.get("published_fps") is not None else "NR"
        lp=r["latency_protocol"]
        rows.append([f"{r['label']} [{ref_index[r['reference_id']]}]",number("map50_pct"),number("map50_95_pct"),number("parameters_m"),number("flops_g"),speed,f"{lp.get('backend') or 'NR'} / {lp.get('precision') or 'dtype NR'}"])
    t("Bảng I Số liệu công bố COCO80 lớp 640px",["Mô hình nguồn","AP50 %","AP95 %","Params M","FLOPs G","Tốc độ T4","Backend dtype"],rows,[2.15,.75,.75,.85,.8,1.2,2.1],True,True)
    p("Bảng I dùng COCO2017 val2017 gộp 80 lớp. GPU tốc độ là T4 nhưng backend, dtype và pipeline không đồng nhất; YOLO11 chỉ ghi TensorRT10, dtype chưa xác định. FPS từ 1000/latency là quy đổi, chưa là phép đo throughput. YOLOv12 và RT-DETRv2 khóa paper v1; không thay bằng Turbo hay checkpoint cập nhật. Baseline tác giả YOLOv13 tự tái lập có AP khác tài liệu chính thức, không ghép nguồn để suy mức cải thiện.")
    t("Bảng II Ma trận khác biệt giao thức",["Thiết lập","Dữ liệu lớp","Accuracy","Latency","Đối sánh"],[
        ["Paper detector","COCO2017; 80 lớp; val2017","640px; AP điểm công bố","T4 TensorRT; dtype tùy nguồn","Bối cảnh; không xếp hạng trực tiếp địa phương"],
        ["Báo cáo trước và đo lại","COCO 10.000 + Adv 1.123; person; test 1.668","416px; val chọn; 5 seed tại mốc riêng","Loại speed cũ và cohort retiming khác bối cảnh","Giữ accuracy và params; không xếp hạng speed n=5"],
        ["Hai nhánh cập nhật","Giữ split; person; test từng xem","416px; 6 epoch; 3 seed ghép","32 warmup + 500 frame; pipeline chung","Trực tiếp hai nhánh mới; thăm dò"]],[1.65,2.2,1.8,1.75,2.1],True)
    p("BadPatch [7] và Adversarial T-shirt [8] cung cấp bối cảnh patch vật lý làm người bị bỏ sót. AP hoặc attack success rate trong bài tấn công khác mục tiêu tăng AP detector. Ảnh patch trong tập dữ liệu chưa chứng minh mô hình chống patch chưa thấy hay giảm DoS tính toán; các kết luận đó cần threat model và thí nghiệm riêng.")
    h("3 Phương pháp và phạm vi đóng góp")
    h("3 1 Cơ sở HyperACE FullPAD và DS C3k2",2)
    p("HyperACE của Lei và cộng sự dùng hyperedge thích nghi để tổng hợp quan hệ nhiều vị trí và thang đo; FullPAD phân phối đặc trưng tới nhiều tầng; DS-C3k2 dùng depthwise separable convolution để giảm phép tính và tham số [1]. Cơ chế tạo tiềm năng linh hoạt về ngữ cảnh, chưa là bằng chứng chống patch. Những kiến trúc và tên thành phần thuộc tác giả YOLOv13.")
    p("Ablation gốc YOLOv13-S đạt AP95 47,1 khi bỏ tunnel và 48,0 khi đủ tunnel; tác giả mô tả bỏ tunnel làm mất tác dụng HyperACE. Đây là tác dụng gộp, chưa là factorial ablation từng phần. Table V gốc cho DS giảm params 11,7 xuống 9,0M và FLOPs 27,1 xuống 20,8G với AP95 cùng làm tròn 48,0 [1]. FLOPs không đo bộ nhớ và kernel; ít FLOPs không bảo đảm nhanh hơn trên Windows PyTorch.")
    h("3 2 Can thiệp vào chuyển giao và recipe",2)
    p("Ứng viên kiểm tra lr thấp, cosine scheduler, AdamW, trọng số box class DFL, mosaic mixup biến đổi hình học và sampling nguồn train. Các yếu tố gộp chưa cho attribution riêng. Ứng viên khôi phục pretrained sửa HyperACE về số hyperedge thực và chuyển hàng person vào classifier một lớp, cùng AdamW cosine; đây là sửa chuyển giao và recipe, không phải phát minh kiến trúc.")
    rows=[]
    for key,cfg in data["protocol"]["candidate_configurations"].items():
        main=f"{cfg.get('optimizer','SGD')}; lr0={cfg.get('lr0')}; lrf={cfg.get('lrf')}; {'cosine' if cfg.get('cos_lr') else 'linear'}; horizon={cfg.get('epochs')}; warmup={cfg.get('warmup_epochs')}"
        extra="; ".join(f"{k}={v}" for k,v in cfg.items() if k not in ("optimizer","lr0","lrf","cos_lr","epochs","warmup_epochs"))
        rows.append([key,main,extra or "Giao thức gốc"])
    t("Bảng III Cấu hình trước sàng lọc",["Cấu hình","Optimizer scheduler","Can thiệp bổ sung"],rows,[1.95,3.75,3.1],True)
    p("Đối chứng giữ horizon 300 nhưng dừng sau 6 epoch; cosine dùng horizon 6. Khác biệt là một can thiệp có chủ đích. Riêng ứng viên adamw_domain_loss_aug dùng sampling có hoàn lại, kỳ vọng 25% lượt ảnh chính thuộc Adv thay khoảng 10,10% train gốc; ảnh phụ mosaic vẫn lấy đều. Sáu epoch có 46.716 lượt chọn ảnh chính mỗi run; không bảo đảm mỗi ảnh xuất hiện sáu lần khi lấy mẫu có hoàn lại. Amendment sửa chuyển giao được khóa trước khi GPU đánh giá ứng viên và khai báo trong protocol.")
    p("Tìm kiếm có điều chỉnh dựa trên validation: bốn cấu hình ban đầu, ứng viên khôi phục chuyển giao sau audit CPU và một ứng viên thứ sáu theo điều kiện đã ghi trước khi đo ứng viên đó. Ứng viên transfer_tiny_lr dùng AdamW 1e-5, đóng băng chín module đầu, augmentation nhẹ và không đóng băng BN. Đây là thử nghiệm thích nghi hữu hạn, chưa là đăng ký trước bất biến từ đầu; không quy tác dụng cho một thành phần riêng.")
    h("4 Thiết lập và kiểm soát chất lượng")
    p("Dữ liệu 11.123 ảnh gồm COCO 10.000 ảnh chọn xác định và mirror Adv 1.123 ảnh khả dụng. Split theo nguồn và ảnh với seed 2026, tỷ lệ 70/15/15: train 7.786, val 1.669, test 1.668; COCO/Adv lần lượt 7.000/786, 1.500/169, 1.500/168. Chỉ lớp person. Byte trùng qua split bằng 0 nhưng chưa loại tương quan người hoặc chuỗi. Có 1.429 ảnh test từ COCO train2017 có thể đã xuất hiện trong pretraining; holdout chỉ độc lập với fine tuning.")
    p("Windows 11, RTX 5060 khoảng 7,96 GiB VRAM, Intel family 6 model 183 với 20 logical CPU, RAM 31,79 GiB. Python chỉ dùng project_yolov13/.venv: Python 3.14.7, Torch 2.11 CUDA 12.8, Ultralytics địa phương 8.3.63. Attention dùng SDPA fallback. Ảnh 416, batch 8, cùng train, sáu epoch; training_config lưu args đầy đủ. Không thay số đo bằng tốc độ TensorRT của bài báo.")
    p(f"Sàng lọc seed 42 có {len(data['screening_records'])} trial chỉ dùng val. Bước 1 chọn epoch có AP95 native lớn nhất trong epoch 1 đến 6, validator AMP FP16 batch16. Bước 2 chỉ đánh giá lại checkpoint đã chọn bằng FP32 batch8 trên toàn validation và xếp hạng ứng viên theo AP95 này. Không đo FP32 mọi epoch, nên có thể bỏ sót checkpoint tối ưu theo FP32. {winner} thắng nhóm cải tiến, khóa tại {data['winner_lock']['locked_at_utc']}; SGD đối chứng không thuộc nhóm thắng. Các nhánh seed 100, 2026, 777 dùng cùng quy tắc chọn checkpoint rồi đánh giá test sau khóa. Test đã được xem trước nên đây là xác nhận thăm dò, chưa là holdout hoàn toàn mới.")
    p("Ba đối chứng SGD dùng đúng sáu epoch đầu của các lượt lịch sử seed100, 2026, 777. Audit seed42 cho thấy toàn bộ giá trị CSV ngoài thời gian và tensor EMA mỗi checkpoint sáu epoch trùng chính xác giữa lần chạy mới và cũ; args có ý nghĩa, SHA nguồn, pretrained và manifest cũng khớp. Chỉ chọn checkpoint trong prefix sáu epoch, không lấy best120 phút. Đối chứng được đo lại test và latency sau khóa trong phiên mới. Chi phí train lịch sử được ghi riêng, chi phí train tăng thêm bằng 0 cho ba prefix tái dùng; không mô tả thành sáu lượt xác nhận đều huấn luyện mới.")
    p("Kết quả trình bày mean ± sample SD và CI 95% Student t trong ngoặc vuông, n=3 và df=2. Δ ghép cùng seed. CI theo optimizer trên một split, không đo bất định dataset, người hoặc quá trình tìm cấu hình. AP95 theo IoU 0,50:0,05:0,95; Precision và Recall theo operating point validator. AP gộp 80 lớp của paper khác AP person.")
    p("Latency mới dùng protocol actual-batch1-txt-loader-v2: danh sách ảnh lưu thành tệp text để LoadImagesAndVideos đọc từng ảnh, thay vì truyền list Python. Callback và forward hook kiểm tra batch thực bằng 1 cho mọi lần dự đoán. Có 532 ảnh test chọn xác định: 32 warmup và 500 frame lấy mẫu, FP32, 416px; median tổng preprocess, inference, postprocess từng seed rồi tổng hợp. Không gồm decode và I/O; P95 lưu nhưng FPS từ median chưa bảo đảm thời gian thực ứng dụng. Params dựa số phần tử tensor; FLOPs từ THOP MACs nhân hai và extrapolation diện tích có thể thiếu custom operators.")
    p("Đính chính phép đo tốc độ: loader nhận list Python trước đây gom 532 ảnh thành một batch lớn dù tham số batch=1, nên mọi latency và FPS lịch sử theo cách đó bị đánh dấu vô hiệu và loại khỏi bảng, đồ thị hoặc kết luận tốc độ cập nhật. Loader cũ còn gây thiếu VRAM ở lần đánh giá đối chứng seed100. Accuracy dùng validator riêng; các số AP đã lưu không bị lỗi loader này. Full test đã tính nhưng chưa lưu trước lần OOM được đánh giá lại, không xem là một record hoàn tất. Mã sửa lưu accuracy trước đo tài nguyên. Speed hiện tại và diagnostic được đo lại với batch thực được kiểm tra; hồ sơ cũ giữ để truy vết, không thay bằng số paper.")
    p("Sau khi sửa batch, cùng checkpoint SGD YOLOv13s seed100 và cùng hash có median 18,973 ms trong phiên xác nhận ghép cặp, nhưng 32,142 ms khi đo lại lịch sử độc lập, lệch khoảng 69,4%. Đây là bất ổn theo bối cảnh runtime chưa được giải thích, không phải cải thiện do kiến trúc hay recipe. Vì vậy toàn bộ latency/FPS cohort lịch sử n=5, kể cả lần đo lại batch1 đúng, bị loại khỏi đối sánh tốc độ trực tiếp; chỉ giữ accuracy và params của chúng. Phân tích speed dùng hai nhánh n=3 trong cohort xác nhận chung; điểm native n=1 chỉ là diagnostic. CI theo seed chưa đo hết sai số giữa các phiên hoặc điều kiện hệ thống.")
    h("5 Kết quả")
    h("5 1 Chẩn đoán và đính chính cấu trúc",2)
    p(f"Pretrained head 80 lớp lọc person đúng tại NMS có validation P {pre['precision']*100:.2f}%, R {pre['recall']*100:.2f}%, AP50 {pre['map50']*100:.2f}%, AP95 {pre['map50_95']*100:.2f}%. Đây là val, không trực tiếp so với test 47,97% trước. Diagnostic 28,54% ban đầu gộp 80 lớp qua single_cls và lọc person sai; record giữ trạng thái invalid_excluded, không dùng làm baseline.")
    p("Audit CPU cho YOLOv13s: pretrained prototype_base thực [8,128], tái dựng YAML [16,128], context_net 1024 lên 2048; 6 tensor HyperACE và 6 classifier khác shape khiến khoảng 1.057.155 phần tử tensor mới không chuyển được. Bản khôi phục 8 edge có 9,024954M params, theo head_transfer_audit.json. Số cũ 9,553338M thuộc biến thể tái dựng thực đo, không gắn trực tiếp checkpoint chính thức. Nano có 4 edge thực và 4 edge tái dựng, không có lỗi HyperACE tương tự; hidden width classifier từ 80 xuống 64 còn khác. Không tổng quát hóa anomaly cho mọi YOLOv13.")
    p("Khôi phục cấu trúc và chuyển hàng person chưa tự xác định nguyên nhân AP thấp. Scheduler, learning rate, head và forgetting cùng ảnh hưởng; cần ablation để quy nguyên nhân. Báo cáo trước và số đo được giữ nguyên, đính chính phạm vi tên mô hình thực chạy.")
    restored=next((r for r in data["screening_records"] if r["configuration"]=="adamw_transfer_restored"),None)
    zero=(restored or {}).get("initialization",{}).get("zero_finetuning_validation")
    if zero:
        p(f"Khởi tạo một lớp đã khôi phục 8 hyperedge và chuyển hàng person, chưa học epoch nào, đạt validation AP50 {zero['map50']*100:.2f}% và AP95 {zero['map50_95']*100:.2f}%. Epoch 0 không tham gia chọn ứng viên đã train. Theo amendment lựa chọn triển khai riêng trước test, checkpoint này được đối chiếu FP32 batch8 với ứng viên đã train thắng sàng lọc; n=1 và không gán CI seed.")
    if deployment.get("choice")=="trained_winner":
        zmap=deployment["zero_validation_fp32"]["map50_95"]
        trained_map=deployment["trained_validation_fp32"]["map50_95"]
        trial=next(r for r in data["screening_records"] if r["configuration"]==winner)
        p(f"Quy tắc triển khai nghiêm ngặt chọn AP95 validation lớn nhất đã ưu tiên nhánh train: {trained_map*100:.4f}% so với khôi phục epoch 0 {zmap*100:.4f}%, tăng {(trained_map-zmap)*100:.4f} pp sau {trial['actual_training_seconds']/60:.2f} phút. Xét độ nhạy Pareto hậu nghiệm với dung sai thực dụng 0,5 pp đã dùng trong nghiên cứu thời gian trước, lợi ích nhỏ có thể làm phương án không train hấp dẫn hơn về chi phí. Phân tích độ nhạy này không thay quyết định khóa và không là tiêu chí chọn ban đầu; cần test mới để kiểm định giá trị ứng dụng.",landscape=True)
    bn=data.get("bn_controlled_diagnostic")
    if bn:
        unchanged=bn['unchanged_fp32_batch8_validation']['map50_95'];reset=bn['bn_restored_fp32_batch8_validation']['map50_95']
        p(f"Kiểm tra cùng FP32 batch8 cho checkpoint adamw_transfer_restored đã train có AP95 {unchanged*100:.4f}%; thay riêng buffer BN bằng giá trị khởi tạo làm AP95 còn {reset*100:.4f}%. Audit xác nhận tensor ngoài BN giữ nguyên. Chênh lệch này không hỗ trợ thao tác đặt lại BN hậu huấn luyện cho checkpoint này, nên ứng viên tiny LR không đóng băng BN theo điều kiện đã khóa. Không suy ra nhân quả toàn quá trình học hoặc mọi mô hình từ một phép hoán đổi buffer. Đánh giá khởi tạo và đã train ở cả FP16 batch16 cũng cho AP95 giảm; sai khác precision đơn thuần không giải thích toàn bộ suy giảm. Mất tri thức pretrained vẫn là diễn giải cần ablation, không phải nguyên nhân đã xác lập.",landscape=True)
    rows=[]
    for r in data["screening_records"]:rows.append([r["configuration"],str(r["selected_epoch"]),f"{r['validation']['map50_95']*100:.2f}",*(f"{r['validation_fp32'][k]*100:.2f}" for k in METRICS),f"{r['actual_training_seconds']/60:.2f}"])
    t("Bảng IV Sàng lọc validation seed42 theo hai bước",["Cấu hình","Epoch chọn","AP95 native %","P FP32 %","R FP32 %","AP50 FP32 %","AP95 FP32 %","Train phút"],rows,[2.7,.65,1.05,.8,.8,1,1,1],True,True)
    p("Sàng lọc một seed không có CI đa seed. Epoch chọn có thể sớm hơn 6 vì lấy validation best. Cột native trích CSV train; FP32 đo lại checkpoint đó, là cơ sở xếp hạng cuối. Cấu hình thắng chỉ tốt nhất trong các ứng viên đã chạy, chưa là tối ưu toàn cục.")
    loss_rows=[]
    for r in data['screening_records']+data['confirmation_records']:
        loss=r['endpoint_loss']
        vector=lambda prefix: '/'.join(f"{loss[prefix+'/'+k+'_loss']:.3f}" for k in ('box','cls','dfl'))
        with project_path(r['results_csv'],must_exist=True).open(encoding='utf-8-sig',newline='') as stream:
            history=[{k.strip():v.strip() for k,v in row.items()} for row in csv.DictReader(stream)]
        start=float(history[0]['metrics/mAP50-95(B)']);end=float(history[-1]['metrics/mAP50-95(B)'])
        loss_rows.append([r['configuration'],str(r['seed']),r['stage'],vector('train'),vector('val'),f'{start*100:.2f} → {end*100:.2f}',f"{r['selected_elapsed_seconds']/60:.2f}"])
    t("Bảng IVa Hội tụ và loss cuối epoch6",["Cấu hình","Seed","Giai đoạn","Train box cls dfl","Val box cls dfl","AP95 native epoch1 → 6 %","Phút tới epoch chọn"],loss_rows,[2.45,.6,1.05,1.4,1.4,1.65,.8],True,True)
    p("Loss là giá trị native cuối quỹ đạo, không phải loss của checkpoint chọn nếu epoch tốt nhất sớm hơn sáu. Train dùng augmentation, val không dùng cùng phép biến đổi, nên khoảng cách loss chưa tự chứng minh overfitting. AP giảm trong khi loss train giảm có thể gợi ý suy giảm chuyển giao hoặc không khớp mục tiêu; sáu epoch và một split chưa phân biệt chắc chắn overfitting, underfitting và forgetting, cũng chưa chứng minh trần hội tụ.")
    h("5 2 Đối chứng mới và ba seed",2)
    accuracy=[];cost=[]
    for g in(control,proposed):
        name="Đối chứng SGD" if g is control else "Nhánh tối ưu";m=g["metrics"]
        accuracy.append([name,"3",*(fmt(m[k],100) for k in METRICS)])
        cost.append([name,fmt(m["parameters_m"],1,3,False),fmt(m["flops_g"],1,2,False),fmt(g["actual_training_seconds"],1/60),fmt(m["latency_ms"]),fmt(m["fps"])])
    accuracy.append(["Native pretrained","1",*(f"{native_test[k]*100:.2f}" for k in METRICS)])
    cost.append(["Native pretrained",f"{native_test['parameters_m']:.3f}",f"{native_test['flops_g']:.2f}","0 fine tune",f"{native_latency['latency_median_ms_per_frame']:.2f}",f"{native_latency['fps_from_median_latency']:.2f}"])
    if zero_test:
        zt=zero_test['test'];zl=zero_test['latency']
        accuracy.append(["Khôi phục epoch0","1",*(f"{zt[k]*100:.2f}" for k in METRICS)])
        cost.append(["Khôi phục epoch0",f"{zt['parameters_m']:.3f}",f"{zt['flops_g']:.2f}","0 fine tune",f"{zl['latency_median_ms_per_frame']:.2f}",f"{zl['fps_from_median_latency']:.2f}"])
    t("Bảng V Accuracy test sau khóa",["Nhánh","n","Precision %","Recall %","AP50 %","AP95 %"],accuracy,[1.65,.35,1.9,1.9,1.9,1.9],True,True)
    t("Bảng VI Tài nguyên và chi phí",["Nhánh","Params M","FLOPs G","Train phút","Latency ms","FPS"],cost,[1.5,.85,.85,2,2,2],True,True)
    p("Hàng native pretrained là một checkpoint cố định head 80 lớp, lọc COCO person trong NMS đúng protocol; n=1 biểu thị một điểm chẩn đoán, không nhóm seed và không gán CI. Đo test và latency sau winner lock. Params/FLOPs khác nhánh một lớp; latency mới cùng batch1 FP32 nhưng dự đoán được lọc person. Chi phí pretraining từ nguồn ngoài không bằng 0; số 0 chỉ là fine tuning bổ sung của dự án.")
    if zero_test:
        role="lựa chọn triển khai đã khóa trên validation" if deployment["choice"]=="zero_restored_person" else "điểm chẩn đoán bổ sung; lựa chọn đã khóa vẫn là nhánh đã huấn luyện được chọn"
        p(f"Hàng khôi phục epoch0 giữ kiến trúc 8 hyperedge và head person một lớp, dùng deterministic weight transfer. Đây là {role}, không là một seed thứ tư của nhánh tối ưu. Không lấy test cao hơn để đổi nhánh sau khóa; CI n=3 và điểm n=1 có đối tượng bất định khác nhau.")
    p(f"Δ tối ưu trừ đối chứng AP95 là {delta_text} pp. {verdict} Nếu nhánh khôi phục cấu trúc thắng, Δ phản ánh gói can thiệp, chưa tách cấu trúc và recipe. Tổng runtime CSV của các trial được phân tích là {sum(r['actual_training_seconds'] for r in data['screening_records']+data['confirmation_records'])/3600:.2f} giờ, chưa cộng khởi tạo và đánh giá; nếu có tái dùng prefix đối chứng lịch sử, tổng này không phải toàn bộ chi phí tính toán mới của vòng cập nhật.")
    fresh_records=data['screening_records']+[r for r in data['confirmation_records'] if r['configuration']!='original_sgd']
    old_seconds=sum(r['actual_training_seconds'] for r in control['seed_records'])
    p(f"Chi phí train CSV của {len(fresh_records)} quỹ đạo mới là {sum(r['actual_training_seconds'] for r in fresh_records)/3600:.2f} giờ; ba prefix đối chứng có thời gian lịch sử cộng {old_seconds/3600:.2f} giờ nhưng không train mới. CSV không tính toàn bộ tải dữ liệu, khởi tạo, đánh giá, diagnostic, thời gian chờ và thời gian gián đoạn phiên làm việc. Ngân sách sáu epoch tương đương lượt dữ liệu, chưa tương đương thời gian đồng hồ hoặc dung lượng mô hình.")
    rows=[]
    for seed in data["protocol"]["confirmation_seeds"]:
        a=next(r for r in control["seed_records"] if r["seed"]==seed);b=next(r for r in proposed["seed_records"] if r["seed"]==seed)
        rows.append([str(seed),str(a["selected_epoch"]),str(b["selected_epoch"]),f"{a['test']['map50_95']*100:.2f}",f"{b['test']['map50_95']*100:.2f}",f"{(b['test']['map50_95']-a['test']['map50_95'])*100:+.2f}"])
    t("Bảng VII Test ghép cặp từng seed",["Seed","Epoch đối chứng","Epoch tối ưu","AP95 đối chứng %","AP95 tối ưu %","Δ pp"],rows,[.8,1.3,1.3,2,2,.9],True,True)
    h("5 3 Phân tầng nguồn",2)
    rows=[]
    for g in(control,proposed):
        for domain in("coco","advtshirt"):
            d=g["domain_metrics"]["domain_test"][domain]
            rows.append(["Đối chứng" if g is control else "Tối ưu",domain,str(d["images"]),str(d["person_instances"]),str(d["negative_images"]),fmt(d["metrics"]["map50"],100),fmt(d["metrics"]["map50_95"],100)])
    t("Bảng VIII Test theo nguồn với ba seed",["Nhánh","Nguồn","Ảnh","Person","Ảnh âm","AP50 %","AP95 %"],rows,[1.25,1,.6,.7,.7,2.25,2.25],True,True)
    p("Pooled AP không bằng mean AP hai nguồn. Test COCO 1.500 ảnh so với Adv 168 nên pooled có thể che miền patch. Phân tầng giúp định vị cải thiện, chưa là causal defense: thiếu cặp cùng cảnh clean và attacked, patch mới và split theo người hoặc chuỗi.")
    h("5 4 Đối chiếu các mốc thời gian đã khảo sát",2)
    prior=read("results/fullscale_experiments/latency_records.json")["sweet_spots"]
    time_rows=[]
    for key in("yolo11n","yolov12n","yolov13n","yolov13s"):
        item=prior[key];curve=item["mean_validation_map50_95_by_budget"]
        time_rows.append([item["label"],*(f"{curve[str(b)]*100:.4f}" if str(b) in curve else "NR" for b in(5,10,20,30,60,120)),str(item["sweet_spot_minutes"])])
    t("Bảng IX Đường cong thời gian thăm dò trước với validation AP95 phần trăm",["Mô hình","5 phút","10 phút","20 phút","30 phút","60 phút","120 phút","Sweet spot phút"],time_rows,[1.4,1,1,1,1,1,1,1.1],True,True)
    p("Nguồn Bảng IX là trường sweet_spots trong results/fullscale_experiments/latency_records.json, chỉ dùng metadata accuracy, loại toàn bộ speed cũ. Mỗi ô là mean validation AP95 của checkpoint tốt nhất đã hoàn tất trong ngân sách với đủ năm seed. NR ở YOLOv13s 5 phút vì seed2026 chưa hoàn tất epoch đầu, nên thiếu nhóm đủ năm seed. Sweet spot là mốc sớm nhất có mean cách peak quan sát không quá 0,5 pp. Đây là các đường cong thăm dò trước, không phải lượt chạy mới 1, 2, 4 hoặc 8 giờ.")
    nano_extra=(prior["yolov13n"]["peak_validation_map50_95"]-prior["yolov13n"]["mean_validation_map50_95_by_budget"]["5"])*100
    p(f"Checkpoint best của YOLOv13s giữ nguyên từ 10 đến 120 phút; YOLOv13n tăng thêm {nano_extra:.4f} pp ở 120 so với 5 phút, nằm trong dung sai 0,5 pp. YOLO11n và YOLOv12n còn tăng tới 120 phút, nên chưa chứng minh đã hội tụ. Đường best-so-far phẳng có thể chỉ do giữ một checkpoint sớm; không đồng nghĩa loss hoặc mô hình ngừng học. 120 phút là mốc tối đa đã thử, chưa có bằng chứng cho kết luận phổ quát 4 hoặc 8 giờ. Cấu trúc tái dựng và recipe của đường cong trước khác nhánh khôi phục hiện tại, vì vậy không chuyển nguyên sweet spot cũ sang mô hình sửa.")
    p("Khuyến nghị khảo sát bắt đầu bằng diagnostic epoch0, các mốc sớm 5, 10, 20 và 30 phút, rồi mở tới 60 hoặc 120 phút khi validation còn tăng vượt dung sai. Với ứng viên tiny LR hiện tại, sáu epoch và diagnostic epoch0 là bằng chứng trực tiếp; cần lập đường cong ngân sách riêng sau sửa chuyển giao trước khi đề xuất mốc dừng cho nghiên cứu tiếp.")
    h("6 Hình ảnh và đánh đổi")
    historical_speed="Accuracy và params lịch sử n=5 được giữ; speed cũ và cohort retiming lệch bối cảnh bị loại."
    points="Native80 lọc person là điểm n=1, không SD/CI."
    if zero_test:points+=" Khôi phục epoch0 là head một lớp 8 hyperedge, cũng n=1 không SD/CI."
    blocks.append(dict(kind="image",path="results/figures/comparison_baseline_vs_proposed.png",landscape=True,caption="Hình 1 Đối sánh AP, FPS và params. "+historical_speed+" Hai nhánh mới n=3 có CI 95% theo seed; speed mới FP32 416px RTX5060. "+points+" Panel COCO80 lớp 640px là bối cảnh khác giao thức. Nguồn: fullscale accuracy, final_summary.json, retiming và benchmark_matrix.json. Ảnh 300 DPI."))
    blocks.append(dict(kind="image",path="results/figures/tradeoff_map_vs_latency.png",landscape=True,caption="Hình 2 AP95 và latency test person, batch1 thực FP32 416px, không decode/I/O. Chỉ hai nhánh ghép cặp n=3 có CI 95% và diagnostic native80 n=1 không CI xuất hiện. Mọi điểm speed lịch sử n=5 bị loại vì loader cũ và drift cohort. Ngưỡng 33,3 ms chỉ tham khảo 30 FPS, không kết luận tăng tốc. Nguồn: final_summary.json và audit timing. Ảnh 300 DPI."))
    p(f"Nhánh tối ưu có AP95 {proposed['metrics']['map50_95']['mean']*100:.2f}% và latency đo lại {proposed['metrics']['latency_ms']['mean']:.2f} ms. Chỉ mô tả chênh lệch speed giữa hai nhánh n=3 trong cùng cohort xác nhận; chưa xác lập tăng tốc do recipe, nhất là khi đã quan sát drift giữa các phiên. Điểm lịch sử n=5 chỉ dùng accuracy và params. Các nhóm khác ngân sách và dung lượng chưa cô lập HyperACE. Cần benchmark lặp theo phiên, kiểm soát trạng thái GPU và export trên thiết bị đích, đo tail p95 cùng decode I/O trước cam kết 30 FPS.")
    h("7 Ưu nhược điểm và giải trình")
    p("Sáu chủ đề giải trình dưới đây do nghiên cứu phân rã từ góp ý tổng quát về AP thấp, tính độc đáo và đối sánh. Chúng không phải sáu câu hỏi nguyên văn hoặc nhận xét cụ thể đã được cung cấp bởi Hội đồng.")
    t("Bảng X Ưu nhược điểm và mức chứng cứ",["Thành phần","Ưu điểm","Hạn chế","Chứng cứ"],[
        ["HyperACE FullPAD","Tương quan đa vị trí đa thang","Backend và dataset ảnh hưởng; chưa chứng minh chống patch","Cơ chế và ablation gốc [1]; chưa ablation địa phương"],
        ["DS C3k2","Giảm params và FLOPs","Không bảo đảm latency","Table V gốc [1]; pipeline đo local"],
        ["Khôi phục pretrained","Giữ shape tương thích","Gộp cùng recipe; chưa causal","Audit tensor và trial mới"],
        ["Scheduler optimizer","Thích ứng horizon","6 khác 300; chưa tối ưu toàn cục","Cùng 6 epoch, seed ghép cặp"],
        ["Loss aug sampling","Ưu tiên miền Adv và điều chỉnh loss","Gộp yếu tố; lặp cảnh có thể overfit","Chỉ toàn cấu hình; cần ablation tiếp"],
        ["Ba seed","Đo biến thiên optimizer","CI rộng, chưa đo bất định dataset","Thăm dò n=3"],
        ["COCO Adv","Cảnh thường và áo patch","Pretrain exposure, tương quan frame","Phân tầng, chưa causal defense"]],[1.5,2.5,2.8,2.3],True,True)
    answers=[
        ("7 1 Vì sao kết quả trước thấp","Diagnostic person đúng cho AP cao hơn, còn head và shape HyperACE khác khi tái dựng; lịch học và forgetting cũng ảnh hưởng. Giữ số trước và khai báo erratum; chưa quy một nguyên nhân khi chưa có ablation."),
        ("7 2 Đóng góp độc đáo","Kiến trúc thuộc YOLOv13 gốc. Đóng góp dự án là audit chuyển giao, recipe person, khóa cấu hình bằng val, seed ghép cặp, phân tầng và đo chi phí. Không nhận kiến trúc gốc là phát minh mới."),
        ("7 3 Tại sao hướng đi tốt hơn",f"Lý thuyết hỗ trợ ngữ cảnh bậc cao. Bằng chứng gói can thiệp là Δ AP95 {delta_text} pp. {verdict} Chưa quy riêng HyperACE, head hoặc scheduler."),
        ("7 4 Có vượt SOTA","Chưa có cơ sở SOTA. COCO80 lớp 640px T4 khác person416 RTX5060. Cần cùng dataset, class mapping, backend và dung lượng mô hình."),
        ("7 5 Diễn giải FLOPs và tốc độ thế nào","Attention kernels, bộ nhớ, depthwise, graph và postprocess tùy backend; FLOPs chưa đo thời gian. Trước hết phải sửa loader để xác nhận batch thực. Chỉ dùng speed mới đã kiểm tra, loại mọi số lịch sử vô hiệu và benchmark export trên thiết bị đích."),
        ("7 6 Thực nghiệm tiếp theo","Ablation khôi phục cấu trúc, head và recipe; tách loss aug sampling; thêm baseline small, seed mới, holdout theo người hoặc chuỗi. Dùng patch mới, clean attacked ghép cặp và ASR nếu kết luận robustness; workload tài nguyên và latency đầu cuối nếu kết luận DoS.")]
    for q,a in answers:h(q,2);p(a)
    h("8 Hạn chế và tái lập")
    p("Val và test tái dùng gây selection bias; khóa mới không tạo test hoàn toàn mới. CI ba seed trên split cố định chưa đo generalization quần thể. Exposure COCO và ảnh tương quan; mirror 1.123 khác 1.131 ảnh công bố. Sáu epoch chưa chứng minh saturation; cosine lịch 6 so với đối chứng lịch 300 dừng sớm. Gói can thiệp cấu trúc, initialization và recipe cần ablation. Lựa chọn hai bước chỉ xếp hạng checkpoint native tốt nhất, chưa tìm tối ưu mọi epoch theo FP32. Ứng viên thứ sáu và đối chiếu epoch0 có thích nghi từ validation; đối chứng dùng prefix lịch sử đã audit, không hoàn toàn chạy đồng thời với nhánh mới.")
    recovery=data.get("interruption_recovery")
    if recovery:
        partial_seconds=recovery.get("known_completed_training_seconds")
        overhead=f" Hai epoch hoàn tất của lượt bỏ dở dùng ít nhất {partial_seconds/60:.2f} phút, chưa gồm phần epoch 3 đang chạy và overhead." if partial_seconds is not None else ""
        p("Phiên thực nghiệm bị gián đoạn giữa ngày 13 và ngày 17/10/2026. Ứng viên tiny LR được chạy lại thành quỹ đạo sáu epoch hoàn chỉnh; hai epoch đầu so với prefix gián đoạn trùng mọi trường CSV ngoài thời gian theo interruption_replay_audit.json. Prefix gián đoạn không tham gia xếp hạng hoặc gộp thêm như seed mới."+overhead+" Runtime các quỹ đạo chọn là phần chi phí xác định được, chưa bao gồm toàn bộ tính toán, diagnostic và đánh giá.")
    p("Checkpoint và log cục bộ, protocol.json, winner_lock.json, final_summary.json, audit tensor và diagnostic invalid được giữ. Hash manifest, pretrained và đầu vào hình cho phép truy vết. Số liệu cũ giữ nguyên. Git chứa artifact compact, không gồm ảnh dataset hoặc checkpoint lớn.")
    h("9 Kết luận")
    p(f"{winner} tốt nhất trong các ứng viên đã train được sàng lọc. Test AP95 {proposed['metrics']['map50_95']['mean']*100:.2f}% so với SGD {control['metrics']['map50_95']['mean']*100:.2f}%; Δ {delta_text} pp. {verdict} Dùng làm điểm khởi đầu vòng xác nhận, lấy split mới và benchmark edge. Kết luận giới hạn trong thiết lập đã đo.")
    if deployment["choice"]=="zero_restored_person" and zero_test:
        p(f"Khuyến nghị triển khai trước mắt theo quyết định validation là giữ checkpoint person khôi phục ở epoch0, test AP95 {zero_test['test']['map50_95']*100:.2f}%, thay vì giả định train lâu hơn sẽ tốt hơn. Chỉ xác nhận lựa chọn cuối trên holdout mới, có phân tầng patch và tốc độ đầu cuối ở thiết bị đích; không suy ra thời gian tối ưu phổ quát từ nghiên cứu sáu epoch.")
    else:
        p(f"Theo quy tắc validation AP95 lớn nhất, lựa chọn triển khai đã khóa là checkpoint đã train của {winner}. Điểm epoch0 nếu được đo test chỉ là đối chiếu chẩn đoán n=1, chưa thay quyết định. Cần xác nhận lựa chọn này trên holdout mới, có phân tầng patch và tốc độ đầu cuối ở thiết bị đích; không suy ra thời gian tối ưu phổ quát từ sáu epoch.")
    native_delta=(proposed["metrics"]["map50_95"]["mean"]-native_test["map50_95"])*100
    p(f"So với điểm chẩn đoán native pretrained AP95 {native_test['map50_95']*100:.2f}%, chênh lệch trung bình nhánh tối ưu là {native_delta:+.2f} pp, chưa có kiểm định đa seed tương ứng cho checkpoint native. Khuyến nghị giữ phương án không fine tune như một baseline thực dụng, và chỉ bỏ nó khi dữ liệu mới cùng chi phí thực tế hỗ trợ lựa chọn khác.")
    h("Tài liệu tham khảo")
    for i,r in enumerate(lit["references"],1):p(f"[{i}] {', '.join(r['authors'])}, “{r['title']},” {r.get('publication_type','')}, {r['year']}. {r['url']} (truy xuất 13/10/2026).")
    h("Phụ lục bằng chứng")
    p("Markdown kèm LaTeX booktabs theo thông lệ IEEE. Word là báo cáo một cột có trang ngang cho bảng và hình, chưa là mẫu camera ready IEEE hai cột.")
    for x in("protocol.json","winner_lock.json","deployment_choice_lock.json","final_summary.json","pretrained_validation.json","pretrained_validation_unfiltered_invalid.json","head_transfer_audit.json","historical_control_reuse_audit.json","bn_controlled_diagnostic.json","latency_protocol_correction.json","historical_latency_retimes.json"):p("Bằng chứng: results/optimization/"+x)
    for x in("results/literature_updated/benchmark_matrix.json","results/literature_updated/source_evidence.json","results/figures/figures_manifest.json","results/fullscale_experiments/latency_records.json","results/updated_research/metric_integrity_audit.json"):p("Bằng chứng: "+x)
    return blocks


def markdown_for(blocks):
    parts=["# "+TITLE];latex=[]
    def esc(x):return str(x).replace("_",r"\_").replace("%",r"\%").replace("&",r"\&").replace("±",r"$\pm$").replace("Δ",r"$\Delta$").replace("\n"," ")
    for b in blocks:
        if b["kind"]=="heading":parts.append("#"*(b["level"]+1)+" "+b["text"])
        elif b["kind"]=="paragraph":parts.append(b["text"])
        elif b["kind"]=="image":parts.append(f"![{b['caption'].split('.')[0]}](figures/{Path(b['path']).name})\n\n"+b["caption"])
        else:
            parts.append(b["caption"]+"\n\n"+"\n".join("| "+" | ".join(str(c).replace("\n","<br>").replace("|","/") for c in row)+" |" for row in [b["headers"],["---"]*len(b["headers"]),*b["rows"]]))
            if b["latex"]:
                lines=[r"\begin{table*}[t]",r"\centering",r"\caption{"+esc(b["caption"])+"}",r"\label{tab:updated-"+str(len(latex)+1)+"}",r"\begin{tabular}{"+"l"+"c"*(len(b["headers"])-1)+"}",r"\toprule"," & ".join(map(esc,b["headers"]))+r" \\",r"\midrule"]
                lines += [" & ".join(map(esc,row))+r" \\" for row in b["rows"]]
                latex.append("\n".join(lines+[r"\bottomrule",r"\end{tabular}",r"\end{table*}"]))
    fence=chr(96)*3
    parts.append("## Mã LaTeX bảng IEEE\n\nDùng gói booktabs. Bảng rộng cần chỉnh nhãn hoặc trang ngang khi tích hợp manuscript.\n\n"+fence+"latex\n"+"\n\n".join(latex)+"\n"+fence)
    return "\n\n".join(parts)+"\n"


def docx_for(blocks):
    from docx import Document
    from docx.enum.section import WD_ORIENT,WD_SECTION_START
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT,WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches,Pt,RGBColor
    doc=Document();doc.core_properties.author=doc.core_properties.last_modified_by="Mai Quốc Bảo";doc.core_properties.title=TITLE;doc.core_properties.language="vi-VN"
    doc.core_properties.created=doc.core_properties.modified=datetime.now(timezone.utc)
    for name,size,bold in(("Normal",12,False),("Title",18,True),("Heading 1",15,True),("Heading 2",13,True),("Caption",12,False)):
        s=doc.styles[name];s.font.name="Times New Roman";s.font.size=Pt(size);s.font.bold=bold;s.font.color.rgb=RGBColor(0,0,0);s.paragraph_format.line_spacing=1.15
        rp=s.element.get_or_add_rPr();rf=rp.rFonts
        if rf is None:rf=OxmlElement("w:rFonts");rp.insert(0,rf)
        for a in("asciiTheme","hAnsiTheme","eastAsiaTheme","cstheme"):rf.attrib.pop(qn("w:"+a),None)
        for a in("ascii","hAnsi","eastAsia","cs"):rf.set(qn("w:"+a),"Times New Roman")
        if s.element.pPr is not None:
            for e in list(s.element.pPr.findall(qn("w:pBdr"))):s.element.pPr.remove(e)
    doc.styles["Normal"].paragraph_format.space_after=Pt(7)
    def configure(s,land):
        s.orientation=WD_ORIENT.LANDSCAPE if land else WD_ORIENT.PORTRAIT
        s.page_width,s.page_height=(Inches(11),Inches(8.5)) if land else (Inches(8.5),Inches(11))
        s.left_margin=s.right_margin=Inches(.65);s.top_margin=s.bottom_margin=Inches(.6)
    configure(doc.sections[0],False);doc.add_paragraph(TITLE,"Title").alignment=1;current=False
    def orient(value):
        nonlocal current
        if value!=current:
            configure(doc.add_section(WD_SECTION_START.NEW_PAGE),value);current=value
            boundary=doc.paragraphs[-1]
            boundary.paragraph_format.line_spacing=Pt(1)
            boundary.paragraph_format.space_before=boundary.paragraph_format.space_after=Pt(0)
    def heading_landscape(index):
        b=blocks[index]
        if b["text"].startswith(("2 Nghiên cứu","7 Ưu nhược")):return True
        following=next((x for x in blocks[index+1:] if x["kind"]!="heading"),None)
        return bool(following and following["kind"] in("table","image") and following.get("landscape"))
    for index,b in enumerate(blocks):
        if b["kind"]=="heading":
            orient(heading_landscape(index));p=doc.add_heading(b["text"],b["level"])
            if b["text"].startswith("7 Ưu nhược"):p.paragraph_format.page_break_before=True
        elif b["kind"]=="paragraph":
            if b.get("landscape") is not None:orient(b["landscape"])
            elif current:
                nxt=next(((j,x) for j,x in enumerate(blocks[index+1:],index+1) if x["kind"] in("table","image","heading")),None)
                if nxt and nxt[1]["kind"]=="heading" and not heading_landscape(nxt[0]):orient(False)
            doc.add_paragraph(b["text"])
        elif b["kind"]=="image":
            orient(b["landscape"])
            if index>0 and blocks[index-1]['kind']=='image':doc.add_page_break()
            p=doc.add_paragraph();p.alignment=1;p.paragraph_format.keep_with_next=True
            p.add_run().add_picture(str(project_path(b["path"],must_exist=True)),width=Inches(7.7))
            caption=doc.add_paragraph(b["caption"],"Caption")
            caption.paragraph_format.space_after=Pt(3)
        else:
            orient(b["landscape"]);p=doc.add_paragraph(b["caption"],"Caption");p.paragraph_format.keep_with_next=True
            table=doc.add_table(rows=1,cols=len(b["headers"]));table.autofit=False;table.alignment=WD_TABLE_ALIGNMENT.CENTER
            for c,x in zip(table.rows[0].cells,b["headers"]):c.text=str(x)
            for values in b["rows"]:
                for c,x in zip(table.add_row().cells,values):c.text=str(x)
            borders=OxmlElement("w:tblBorders")
            for side in("top","left","bottom","right","insideH","insideV"):
                e=OxmlElement("w:"+side)
                for k,v in(("val","single"),("sz","4"),("color","D9D9D9")):e.set(qn("w:"+k),v)
                borders.append(e)
            table._tbl.tblPr.append(borders)
            for col,width in zip(table.columns,b["widths"]):col.width=Inches(width)
            for i,row in enumerate(table.rows):
                row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
                if i==0:e=OxmlElement("w:tblHeader");e.set(qn("w:val"),"true");row._tr.get_or_add_trPr().append(e)
                for j,(c,width) in enumerate(zip(row.cells,b["widths"])):
                    c.width=Inches(width);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    shd=OxmlElement("w:shd");shd.set(qn("w:fill"),"D9EAF7" if i==0 else "F3F6FA" if i%2==0 else "FFFFFF");c._tc.get_or_add_tcPr().append(shd)
                    mar=OxmlElement("w:tcMar")
                    for side in("top","left","bottom","right"):
                        e=OxmlElement("w:"+side);e.set(qn("w:w"),"80");e.set(qn("w:type"),"dxa");mar.append(e)
                    c._tc.get_or_add_tcPr().append(mar)
                    for p in c.paragraphs:
                        p.alignment=0 if j==0 or len(b["headers"])<=5 else 1
                        p.paragraph_format.line_spacing=1.15;p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3);p.paragraph_format.keep_with_next=(i<len(table.rows)-1 if len(table.rows)<=5 else i==0)
                        for r in p.runs:r.font.name="Times New Roman";r.font.size=Pt(12);r.bold=i==0
    foot=doc.sections[0].footer.paragraphs[0];foot.alignment=1;foot.add_run("Báo cáo nghiên cứu YOLOv13 · ")
    field=OxmlElement("w:fldSimple");field.set(qn("w:instr"),"PAGE");foot._p.append(field)
    for run in foot.runs:run.font.name="Times New Roman";run.font.size=Pt(9)
    target=OUT/(STEM+".docx");doc.save(target);return target


def main():
    require_local_venv();configure_project_environment()
    ap=argparse.ArgumentParser();ap.add_argument("--markdown-only",action="store_true");args=ap.parse_args()
    data,lit,figures=require_inputs();blocks=blocks_for(data,lit,figures);EDITOR.mkdir(parents=True,exist_ok=True);write_json(EDITOR/"editor_content_blocks.json",blocks)
    md=OUT/(STEM+".md");md.write_text(markdown_for(blocks),encoding="utf-8");paths=[md]
    if not args.markdown_only:paths.append(docx_for(blocks))
    write_json(EDITOR/"editor_artifact_manifest.json",{"status":"authored_pending_visual_qa","created_at_utc":datetime.now(timezone.utc).isoformat(),"author":"Mai Quốc Bảo","winner_configuration":data["winner_configuration"],"artifacts":[{"path":str(x.relative_to(project_path("."))),"sha256":sha256_file(x),"bytes":x.stat().st_size} for x in paths],"source_sha256":{x:sha256_file(x) for x in("results/optimization/final_summary.json","results/literature_updated/benchmark_matrix.json","results/figures/figures_manifest.json","results/fullscale_experiments/latency_records.json")}})
    print(json.dumps({"created":[str(x) for x in paths],"blocks":len(blocks)},ensure_ascii=False))


if __name__=="__main__":main()
